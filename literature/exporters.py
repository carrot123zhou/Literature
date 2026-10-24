"""内置多格式导出（纯渲染层，无全局状态）。

数据源只有两个：``document/document.json``（规范文档图谱）与
``translations/<locale>/translation.json``（按 block_id 挂接的译文层）。
本模块只读取它们，绝不修改 raw / document / translations；产物一律写到
项目目录下的 ``exports/<export_id>/``。

格式：
- ``epub``：EPUB3 单文件；页下注块（``type=="footnote"``）能匹配到正文 ``^{N}``
  上标标记时，转成 EPUB3 noteref/aside 真脚注（阅读器弹层）。
- ``docx``：Word 单文件（需要 python-docx）；匹配到的页下注转成 Word 真脚注
  （``word/footnotes.xml`` + 上标引用），未匹配到的保持普通段落。
- ``markdown``：单文件 ``<名称>.md``（与 EPUB/DOCX 同名）+ 同级 ``assets/``。
- ``html``：主入口 ``<名称>.html`` + 分章节 ``chapter-0001.html`` … + ``style.css`` + ``assets/``。

模式：``source``（原文）/ ``translation``（译文，缺译回退原文）/
``bilingual``（原文段后紧跟译文段，缺译只出原文段，不重复）。
"""

from __future__ import annotations

import html
import json
import mimetypes
import re
import shutil
import zipfile
from bisect import bisect_right
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Iterator, Literal
from xml.etree import ElementTree as ET

from .footnote_refs import (
    FootnotePlan,
    SUPERSCRIPT_RE,
    annotate as annotate_footnote_markers,
    build_plan,
    superscript_content,
)
from .models import Block, Document, TranslationEntry, TranslationLayer, utc_now
from .storage import ProjectStore, StorageError, atomic_json, safe_locale, safe_name


ExportMode = Literal["source", "translation", "bilingual"]

FORMATS: list[dict[str, Any]] = [
    {"id": "epub", "label": "EPUB", "single_file": True},
    {"id": "docx", "label": "Word DOCX", "single_file": True},
    {"id": "markdown", "label": "Markdown", "single_file": True},
    {"id": "html", "label": "分页 HTML", "single_file": False},
]
MODES: list[dict[str, str]] = [
    {"id": "source", "label": "原文"},
    {"id": "translation", "label": "译文"},
    {"id": "bilingual", "label": "对照版"},
]
FORMAT_IDS = {item["id"] for item in FORMATS}
MODE_IDS = {item["id"] for item in MODES}

MEDIA_TYPES = {
    ".epub": "application/epub+zip",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".md": "text/markdown; charset=utf-8",
    ".html": "text/html; charset=utf-8",
    ".css": "text/css; charset=utf-8",
    ".json": "application/json; charset=utf-8",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".svg": "image/svg+xml",
}

_EXPORT_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,80}$")

BASE_CSS = """
:root{color:#20322b;background:#f6f5ef;font-family:-apple-system,BlinkMacSystemFont,"Noto Sans CJK SC","Songti SC",sans-serif}
body{max-width:50rem;margin:3rem auto;padding:0 1.5rem;line-height:1.72;color:#20322b}
a{color:#2f5142}
h1,h2,h3,h4,h5,h6{line-height:1.3}
/* 浏览器阅读页：译文用更弱的墨色区分，不再用左侧竖线 */
.translation{color:#52625a}
.source+.translation{margin-top:.35rem}
p{margin:.7rem 0}
img{max-width:100%;height:auto}
table{border-collapse:collapse;width:100%;margin:1rem 0}
td,th{border:1px solid #bac5bd;padding:.45rem;text-align:left;vertical-align:top}
figcaption{color:#63716b;font-size:.92em}
.block{margin:.7rem 0}
pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#eeeee6;padding:.6rem;border-radius:.3rem}
.footnote{color:#63716b;font-size:.94em}
a.noteref{color:inherit;text-decoration:none;font-size:.72em;vertical-align:super}
aside.footnotes{margin-top:1.5rem;padding-top:.9rem;border-top:1px solid #bac5bd}
p.noteref-back{margin:.5rem 0 0;font-size:.85em}
p.noteref-back a{text-decoration:none;color:#2f5142}
""".strip() + "\n"

_INDEX_CSS_EXTRA = """
.toc{background:#fff;border:1px solid #d7ddd7;border-radius:.6rem;padding:1rem 1.4rem}
.toc ol{margin:.2rem 0 .2rem 1.1rem;padding:0}
.toc li{margin:.25rem 0}
.meta{color:#63716b}
""".strip()


# EPUB 专用排版（book.css）：不设背景色/前景色（不干扰阅读器夜间模式），
# 全部用相对单位；serif 正文 + 两端对齐 + 首行缩进，对齐 Apple Books 商店书的排版惯例。
EPUB_CSS = """
body{
  margin:0;
  line-height:1.65;
  font-family:Georgia,"Palatino","Times New Roman","Songti SC","Noto Serif CJK SC","Source Han Serif SC",serif;
  text-align:justify;
  widows:2;
  orphans:2;
}

h1,h2,h3,h4,h5,h6{
  font-family:"Avenir Next","Helvetica Neue",Helvetica,"PingFang SC","Noto Sans CJK SC","Source Han Sans SC",sans-serif;
  font-weight:700;
  line-height:1.3;
  text-align:left;
  hyphens:none;
  -webkit-hyphens:none;
  page-break-after:avoid;
  break-after:avoid;
}

/* 章节内小标题 */
h1{font-size:1.7em;margin:2em 0 .8em}
h2{font-size:1.4em;margin:1.8em 0 .7em}
h3{font-size:1.2em;margin:1.5em 0 .6em}
h4,h5,h6{font-size:1.05em;margin:1.2em 0 .5em}

/* 标题的译文：弱化为常规字重副题（原文标题之下） */
.block-title .translation{
  font-size:.55em;
  font-weight:400;
  margin:.55em 0 0;
  line-height:1.4;
}

/* 章首标题（每章第一个标题块） */
body > section.block-title:first-child{margin:2.2em 0 1.4em}
body > section.block-title:first-child h1,
body > section.block-title:first-child h2{font-size:1.7em;margin:0}
body > section.block-title:first-child h3,
body > section.block-title:first-child h4,
body > section.block-title:first-child h5,
body > section.block-title:first-child h6{font-size:1.45em;margin:0}

p{
  margin:0 0 .5em;
  text-indent:1.5em;
  hyphens:auto;
  -webkit-hyphens:auto;
}
/* 译文以中文为主，按中文排版惯例缩进两字 */
p.translation{text-indent:2em}

/* 标题后的第一段不缩进（出版惯例） */
section.block-title + section.block-paragraph > p:first-child{text-indent:0}

/* 封面页 */
body.cover-page{margin:0;text-align:center}
body.cover-page .cover{margin:0;padding:0}
body.cover-page img{max-width:100%;height:auto}

.footnote{font-size:.88em;margin:.35em 0 .5em;text-indent:0}

ul,ol{margin:.6em 0;padding-left:1.5em}
li{margin:.28em 0;text-align:justify;hyphens:auto;-webkit-hyphens:auto}

pre,code{font-family:"SF Mono",Menlo,Consolas,"Courier New",monospace}
pre{
  font-size:.88em;
  white-space:pre-wrap;
  overflow-wrap:anywhere;
  margin:.8em 0;
  text-indent:0;
  text-align:left;
  page-break-inside:avoid;
  break-inside:avoid;
}

img{max-width:100%;height:auto}
.block-image{text-align:center;margin:1em 0}
figcaption{font-size:.88em;margin:.4em 0 1em;text-indent:0;text-align:center;hyphens:none;-webkit-hyphens:none}

table{
  border-collapse:collapse;
  margin:.9em auto;
  font-size:.92em;
  page-break-inside:avoid;
  break-inside:avoid;
}
td,th{border:1px solid;padding:.4em .55em;text-align:left;vertical-align:top}

a{color:inherit}

/* 页下注转真脚注：正文上标引用 + 章末 aside 注文（阅读器弹层或章末注列表） */
a.noteref{color:inherit;text-decoration:none;font-size:.72em;line-height:0;vertical-align:super}
aside.footnotes{margin-top:1.5em;padding-top:.9em;border-top:1px solid}
p.noteref-back{margin:.5em 0 0;text-indent:0;font-size:.85em}
p.noteref-back a{text-decoration:none}

/* 目录页 */
nav > h1{margin:2em 0 1em}
nav ol{list-style:none;padding-left:0;margin:0}
nav ol ol{padding-left:1.4em}
nav li{margin:.5em 0}
""".strip() + "\n"


# ---------------------------------------------------------------------------
# 译文/内容取值
# ---------------------------------------------------------------------------

def entry_for(layer: TranslationLayer | None, block: Block) -> TranslationEntry | None:
    return layer.entries.get(block.block_id) if layer else None


def has_translation(block: Block, entry: TranslationEntry | None) -> bool:
    """该 block 在当前译文层里是否真的有一段可用译文。"""
    if entry is None:
        return False
    if block.type == "list":
        return bool(entry.items)
    if block.type == "table":
        return bool(entry.cells)
    if block.type == "image":
        return bool(entry.caption)
    return bool((entry.text or "").strip())


def _text_roles(block: Block, entry: TranslationEntry | None, mode: ExportMode) -> list[tuple[str, str]]:
    source = block.source_text or ""
    translated = entry.text if entry and entry.text else ""
    if mode == "source":
        return [("source", source)]
    if mode == "translation":
        return [("translation", translated or source)]
    roles = [("source", source)]
    if translated:
        roles.append(("translation", translated))
    return roles


def _list_roles(block: Block, entry: TranslationEntry | None, mode: ExportMode) -> list[tuple[str, list[str]]]:
    source = list(block.items or [])
    translated = list(entry.items) if entry and entry.items else []
    if mode == "source":
        return [("source", source)]
    if mode == "translation":
        return [("translation", translated or source)]
    roles = [("source", source)]
    if translated:
        roles.append(("translation", translated))
    return roles


def _caption_roles(block: Block, entry: TranslationEntry | None, mode: ExportMode) -> list[tuple[str, list[str]]]:
    source = list(block.caption or [])
    translated = list(entry.caption) if entry and entry.caption else []
    if mode == "source":
        return [("source", source)]
    if mode == "translation":
        return [("translation", translated or source)]
    roles = [("source", source)]
    if translated:
        roles.append(("translation", translated))
    return roles


def _table_roles(block: Block, entry: TranslationEntry | None, mode: ExportMode) -> list[str]:
    if mode == "source":
        return ["source"]
    if mode == "translation":
        return ["translation"]
    roles = ["source"]
    if entry and entry.cells:
        roles.append("translation")
    return roles


def _table_matrix(block: Block, entry: TranslationEntry | None, role: str) -> list[list[str]]:
    if not block.cells:
        return []
    rows = max(cell.row for cell in block.cells) + 1
    columns = max(cell.column for cell in block.cells) + 1
    values = entry.cells if entry and entry.cells and role == "translation" else {}
    matrix = [["" for _ in range(columns)] for _ in range(rows)]
    for cell in block.cells:
        row = max(0, min(rows - 1, cell.row))
        column = max(0, min(columns - 1, cell.column))
        matrix[row][column] = values.get(cell.cell_id, cell.text)
    return matrix


def _table_html(block: Block, entry: TranslationEntry | None, role: str) -> str:
    if not block.cells:
        return str(block.format.get("table_html", "") or "")
    values = entry.cells if entry and entry.cells and role == "translation" else {}
    rows = max(cell.row for cell in block.cells) + 1
    by_row: dict[int, list[Any]] = {}
    for cell in block.cells:
        by_row.setdefault(cell.row, []).append(cell)
    output = ["<table>"]
    for row in range(rows):
        output.append("<tr>")
        for cell in sorted(by_row.get(row, []), key=lambda item: item.column):
            tag = "th" if cell.header else "td"
            attributes = ""
            if cell.rowspan > 1:
                attributes += f' rowspan="{int(cell.rowspan)}"'
            if cell.colspan > 1:
                attributes += f' colspan="{int(cell.colspan)}"'
            output.append(f"<{tag}{attributes}>{html.escape(values.get(cell.cell_id, cell.text))}</{tag}>")
        output.append("</tr>")
    output.append("</table>")
    return "".join(output)


# ---------------------------------------------------------------------------
# 章节选择与切分
# ---------------------------------------------------------------------------

@dataclass(slots=True)
class Section:
    title: str
    level: int
    blocks: list[Block]
    nav: bool = True


def flatten_toc(nodes: Iterable[Any]) -> Iterator[Any]:
    for node in nodes:
        yield node
        yield from flatten_toc(getattr(node, "children", []) or [])


def toc_levels(document: Document) -> dict[str, int]:
    """TOC 节点 block_id → 规范化 level（1..6）。"""
    levels: dict[str, int] = {}
    for node in flatten_toc(document.toc):
        levels[str(node.block_id)] = max(1, min(6, int(node.level or 1)))
    return levels


def _title_positions(blocks: list[Block]) -> dict[int, list[int]]:
    """by_level[L] = 所有「type==title 且 level<=L」block 的下标（升序）。"""
    by_level: dict[int, list[int]] = {level: [] for level in range(1, 7)}
    for index, block in enumerate(blocks):
        if block.type != "title":
            continue
        level = max(1, min(6, int(block.level or 1)))
        for container in range(level, 7):
            by_level[container].append(index)
    return by_level


def _chapter_end(by_level: dict[int, list[int]], start: int, level: int, total: int) -> int:
    positions = by_level[max(1, min(6, int(level or 1)))]
    index = bisect_right(positions, start)
    return positions[index] if index < len(positions) else total


def _block_positions(blocks: list[Block]) -> dict[str, int]:
    positions: dict[str, int] = {}
    for index, block in enumerate(blocks):
        positions.setdefault(block.block_id, index)
    return positions


def select_blocks(
    document: Document,
    chapters: Iterable[str] | None = None,
    block_ids: Iterable[str] | None = None,
) -> list[Block]:
    """按冻结契约展开所选内容。

    - ``block_ids`` 非空时优先：精确取这些 block（保持阅读顺序）。
    - 否则 ``chapters`` 为 TOC 节点 block_id 列表：从该 title block 开始向后收集，
      遇到下一个 ``type=="title" && level<=该节点 level`` 的 block 时停止（不含它），
      起始 title 自身包含在内；空列表 = 全书。
    """
    blocks = list(document.blocks)
    positions = _block_positions(blocks)

    chosen_ids = [str(value) for value in (block_ids or [])]
    if chosen_ids:
        return [blocks[index] for index in sorted({positions[value] for value in chosen_ids if value in positions})]

    chapter_values = [str(value) for value in (chapters or [])]
    if not chapter_values:
        return blocks

    levels = toc_levels(document)
    by_level = _title_positions(blocks)
    ranges: list[tuple[int, int]] = []
    for value in chapter_values:
        start = positions.get(value)
        if start is None:
            continue
        block = blocks[start]
        level = levels.get(value) or max(1, min(6, int(block.level or 1)))
        ranges.append((start, _chapter_end(by_level, start, level, len(blocks))))
    ranges.sort()
    merged: list[tuple[int, int]] = []
    for start, end in ranges:
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    return [block for start, end in merged for block in blocks[start:end]]


def split_sections(document: Document, blocks: list[Block]) -> list[Section]:
    """把已选 block 切成若干个「章节」，供 EPUB / 分页 HTML 的分页与目录使用。

    优先按 TOC 节点切分（保留层级）；无可用 TOC 节点时回退到 ``level<=2`` 的 title。
    """
    if not blocks:
        return []
    positions = _block_positions(blocks)
    starts: list[tuple[int, str, int]] = []
    for node in flatten_toc(document.toc):
        block_id = str(node.block_id)
        if block_id not in positions:
            continue
        index = positions[block_id]
        title = str(node.title or "") or blocks[index].source_text or ""
        starts.append((index, title, max(1, min(6, int(node.level or 1)))))
    if not starts:
        for index, block in enumerate(blocks):
            if block.type == "title" and 1 <= int(block.level or 0) <= 2:
                starts.append((index, block.source_text or "", max(1, min(6, int(block.level or 1)))))
    starts.sort()
    unique: list[tuple[int, str, int]] = []
    seen: set[int] = set()
    for item in starts:
        if item[0] in seen:
            continue
        seen.add(item[0])
        unique.append(item)

    sections: list[Section] = []
    if not unique or unique[0][0] > 0:
        end = unique[0][0] if unique else len(blocks)
        sections.append(Section(document.title, 1, blocks[:end], nav=False))
    for index, (start, title, level) in enumerate(unique):
        end = unique[index + 1][0] if index + 1 < len(unique) else len(blocks)
        sections.append(Section(title, level, blocks[start:end], nav=True))
    return sections


def _section_nav_title(section: Section, layer: TranslationLayer | None, mode: ExportMode) -> str:
    """导航条目标题：译文模式用纯译文标题；对照模式拼「原文 · 译文」；缺失回退原文。"""
    if not section.nav or not section.blocks:
        return section.title
    first = section.blocks[0]
    if first.type != "title":
        return section.title
    entry = entry_for(layer, first)
    translated = (entry.text or "").strip() if entry else ""
    if not translated or translated == (first.source_text or "").strip():
        return section.title
    if mode == "translation":
        return translated
    if mode == "bilingual":
        return f"{section.title} · {translated}"
    return section.title


def _build_nav_tree(items: list[tuple[str, str, int]]) -> list[dict[str, Any]]:
    root: dict[str, Any] = {"level": 0, "children": []}
    stack = [root]
    for title, href, level in items:
        node: dict[str, Any] = {"title": title, "href": href, "level": level, "children": []}
        while len(stack) > 1 and stack[-1]["level"] >= level:
            stack.pop()
        stack[-1]["children"].append(node)
        stack.append(node)
    return root["children"]


def _render_nav(nodes: list[dict[str, Any]]) -> str:
    if not nodes:
        return ""
    parts = ["<ol>"]
    for node in nodes:
        parts.append(f'<li><a href="{html.escape(node["href"])}">{html.escape(_plain_superscripts(node["title"]))}</a>')
        parts.append(_render_nav(node["children"]))
        parts.append("</li>")
    parts.append("</ol>")
    return "".join(parts)


# ---------------------------------------------------------------------------
# 资源（图片）映射与复制
# ---------------------------------------------------------------------------

def asset_plan(document: Document, project: Path) -> dict[str, str]:
    """resource_id → 输出文件名（basename 冲突用 ``<resource_id>-`` 前缀去重）。"""
    mapping: dict[str, str] = {}
    used: dict[str, str] = {}
    for resource_id, value in document.resources.items():
        relative = str((value or {}).get("path") or "")
        if not relative:
            continue
        source = project / relative
        if not source.is_file():
            continue
        base = Path(relative).name
        candidate = base
        if candidate in used:
            candidate = f"{resource_id}-{base}"
        suffix = 2
        while candidate in used:
            candidate = f"{resource_id}-{suffix}-{base}"
            suffix += 1
        used[candidate] = resource_id
        mapping[resource_id] = candidate
    return mapping


def _copy_assets(project: Path, target: Path, document: Document, mapping: dict[str, str]) -> list[Path]:
    files: list[Path] = []
    if not mapping:
        return files
    target.mkdir(parents=True, exist_ok=True)
    for resource_id, filename in mapping.items():
        relative = str((document.resources.get(resource_id) or {}).get("path") or "")
        if not relative:
            continue
        source = project / relative
        if not source.is_file():
            continue
        destination = target / filename
        if not destination.exists():
            shutil.copy2(source, destination)
        files.append(destination)
    return files


# ---------------------------------------------------------------------------
# Markdown 渲染
# ---------------------------------------------------------------------------

def _markdown_block(block: Block, entry: TranslationEntry | None, mode: ExportMode, assets: dict[str, str]) -> list[str]:
    lines: list[str] = []
    kind = block.type
    if kind == "title":
        level = max(1, min(6, int(block.level or 1)))
        for role, text in _text_roles(block, entry, mode):
            if not text:
                continue
            lines.extend([f"{'#' * level} {text}", ""])
    elif kind in {"paragraph", "code", "footnote"}:
        for role, text in _text_roles(block, entry, mode):
            if not text:
                continue
            prefix = "> [页下注] " if kind == "footnote" else ""
            if mode == "bilingual" and role == "translation":
                prefix += "> "
            lines.extend([prefix + text, ""])
    elif kind == "list":
        for role, values in _list_roles(block, entry, mode):
            for value in values:
                prefix = "> " if (mode == "bilingual" and role == "translation") else ""
                lines.append(f"{prefix}- {value}")
            lines.append("")
    elif kind == "equation":
        lines.extend(["$$", block.source_text, "$$", ""])
    elif kind == "table":
        for role in _table_roles(block, entry, mode):
            matrix = _table_matrix(block, entry, role)
            if not matrix:
                fallback = str(block.format.get("table_html", "") or "")
                if fallback:
                    lines.extend([fallback, ""])
                continue
            lines.append("| " + " | ".join(matrix[0]) + " |")
            lines.append("| " + " | ".join("---" for _ in matrix[0]) + " |")
            lines.extend("| " + " | ".join(row) + " |" for row in matrix[1:])
            lines.append("")
    elif kind == "image":
        for resource_id in block.resource_refs:
            filename = assets.get(resource_id)
            if filename:
                lines.append(f"![](assets/{filename})")
        for role, captions in _caption_roles(block, entry, mode):
            for caption in captions:
                prefix = "> " if (mode == "bilingual" and role == "translation") else ""
                lines.append(prefix + caption)
        lines.append("")
    else:
        if block.source_text:
            lines.extend([block.source_text, ""])
    return lines


def render_markdown(
    document: Document, blocks: list[Block], layer: TranslationLayer | None, mode: ExportMode,
    assets: dict[str, str] | None = None, plan: FootnotePlan | None = None,
) -> str:
    mapping = assets if assets is not None else asset_plan(document, Path("."))
    lines: list[str] = []
    for block in blocks:
        if plan is not None and block.block_id in plan.absorbed:
            continue  # 续行块已并入前注，不重复输出
        lines.extend(_markdown_block(block, entry_for(layer, block), mode, mapping))
    text = "\n".join(lines).rstrip() + "\n"
    for resource_id, filename in mapping.items():
        text = text.replace(f"({resource_id})", f"(assets/{filename})")
    return text


# ---------------------------------------------------------------------------
# HTML / XHTML 渲染
# ---------------------------------------------------------------------------

def _lang_attr(role: str, mode: ExportMode, source_language: str) -> str:
    """对照模式下给原文标 lang，让西文在两端对齐时能正确连字断行。"""
    if mode == "bilingual" and role == "source" and source_language:
        value = html.escape(source_language)
        return f' xml:lang="{value}" lang="{value}"'
    return ""


class _RefContext:
    """一次导出内的引用渲染状态：noteref 序号、注文所在文件、当前输出文件。

    同一注文被多章同号标记共享时，aside 只在首锚章节渲染一次；其他章节的
    noteref 需要跨文件链接（``text/section-0007.xhtml#fn-…``）。
    """

    def __init__(self) -> None:
        self.counters: dict[str, int] = {}
        self.note_file: dict[str, str] = {}
        self.current: str = ""


def _epub_note_anchor(note_block_id: str) -> str:
    return "fn-" + re.sub(r"[^A-Za-z0-9_-]", "_", note_block_id)


def _epub_noteref(number: str, note_block_id: str, ctx: "_RefContext | None" = None) -> str:
    anchor = _epub_note_anchor(note_block_id)
    if ctx is None:
        return f'<a class="noteref" epub:type="noteref" href="#{anchor}"><sup>{html.escape(number)}</sup></a>'
    ordinal = ctx.counters.get(note_block_id, 0) + 1
    ctx.counters[note_block_id] = ordinal
    target = ctx.note_file.get(note_block_id) or ""
    href = f"#{anchor}" if not target or target == ctx.current else f"{target}#{anchor}"
    return (f'<a id="ref-{anchor}-{ordinal}" class="noteref" epub:type="noteref" '
            f'href="{href}"><sup>{html.escape(number)}</sup></a>')


def _escape_with_refs(text: str, plan: FootnotePlan | None, block: Block, role: str,
                      by_offset: bool | None = None,
                      ctx: "_RefContext | None" = None) -> str:
    """HTML 转义；已绑定的 ``^{N}`` 标记换成 noteref，其余 ``^{…}`` 按上标渲染。"""
    if not text:
        return ""
    if plan is None:
        return _html_text(text)
    segments = annotate_footnote_markers(text, plan, block, role, by_offset)
    return _render_html_segments(segments, ctx)


def _html_text(value: str) -> str:
    """转义文本，并把剩余的 ``^{…}`` 渲染为 ``<sup>``。"""
    parts: list[str] = []
    cursor = 0
    for match in SUPERSCRIPT_RE.finditer(value):
        if match.start() > cursor:
            parts.append(html.escape(value[cursor:match.start()]))
        parts.append(f"<sup>{html.escape(superscript_content(match.group(1)))}</sup>")
        cursor = match.end()
    if not parts:
        return html.escape(value)
    if cursor < len(value):
        parts.append(html.escape(value[cursor:]))
    return "".join(parts)


def _render_html_segments(segments: list[tuple[str, str, str]],
                          ctx: "_RefContext | None" = None) -> str:
    parts: list[str] = []
    for kind, value, note_id in segments:
        if kind == "ref":
            parts.append(_epub_noteref(value, note_id, ctx))
        else:
            parts.append(_html_text(value))
    return "".join(parts)


def _drop_leading_note_marker(segments: list[tuple[str, str, str]], number: str) -> list[tuple[str, str, str]]:
    """脚注正文自身的引导编号（``^{N}``）不渲染成引用，也不留在注文里。"""
    if not segments or segments[0][0] != "text" or not number:
        return segments
    prefix = f"^{{{number}}}"
    value = segments[0][1]
    if not value.startswith(prefix):
        return segments
    rest = value[len(prefix):]
    if rest.strip():
        return [("text", rest.lstrip(), "")] + segments[1:]
    return segments[1:]


def _note_role_texts(link: "NoteLink", plan: FootnotePlan, layer: TranslationLayer | None,
                     mode: ExportMode) -> list[tuple[str, str]]:
    """注文的角色文本；续行块并入后，其译文也拼进同一条注文。"""
    source = link.text
    translated_parts: list[str] = []
    entry = entry_for(layer, link.block)
    if entry is not None and entry.text:
        translated_parts.append(entry.text)
    for absorbed_id, target_id in plan.absorbed.items():
        if target_id != link.block_id:
            continue
        absorbed_entry = entry_for(layer, plan.absorbed_blocks[absorbed_id])
        if absorbed_entry is not None and absorbed_entry.text:
            translated_parts.append(absorbed_entry.text)
    translated = " ".join(part.strip() for part in translated_parts if part and part.strip())
    if mode == "source":
        return [("source", source)]
    if mode == "translation":
        return [("translation", translated or source)]
    roles = [("source", source)]
    if translated:
        roles.append(("translation", translated))
    return roles


def _epub_footnote_asides(note_ids: list[str], plan: FootnotePlan, layer: TranslationLayer | None,
                          mode: ExportMode, source_language: str,
                          ctx: "_RefContext | None" = None) -> str:
    """把一组页下注渲染为章节末尾的 EPUB3 footnote aside（含 ↩ 回链）。"""
    parts: list[str] = []
    for note_id in note_ids:
        link = plan.links.get(note_id)
        if link is None:
            continue
        block = link.block
        anchor = _epub_note_anchor(note_id)
        inner = [
            f'<section class="block block-footnote" data-block-id="{html.escape(block.block_id)}"'
            f' data-page="{block.page_index + 1}">'
        ]
        for role, text in _note_role_texts(link, plan, layer, mode):
            if not text:
                continue
            segments = _drop_leading_note_marker([("text", text, "")], link.number)
            rendered = _render_html_segments(segments, ctx)
            if not rendered:
                continue
            inner.append(
                f'<p class="footnote {role}"{_lang_attr(role, mode, source_language)}>{rendered}</p>'
            )
        inner.append("</section>")
        back = f"#ref-{anchor}-1"
        inner.append(f'<p class="noteref-back"><a epub:type="footnote-return" href="{back}">↩</a></p>')
        parts.append(f'<aside class="footnotes" epub:type="footnote" id="{anchor}">{"".join(inner)}</aside>')
    return "".join(parts)


def _block_html(block: Block, entry: TranslationEntry | None, mode: ExportMode, assets: dict[str, str],
                source_language: str = "", plan: FootnotePlan | None = None,
                ctx: "_RefContext | None" = None) -> str:
    kind = block.type
    if kind == "footnote" and plan is not None and (
            block.block_id in plan.links or block.block_id in plan.absorbed):
        return ""  # 已绑定注文移到章末 aside（续行块已并入前注），正文流不再重复
    output = [
        f'<section class="block block-{html.escape(kind)}"'
        f' data-block-id="{html.escape(block.block_id)}" data-page="{block.page_index + 1}">'
    ]
    if kind == "title":
        tag = f"h{max(1, min(6, int(block.level or 1)))}"
        for role, text in _text_roles(block, entry, mode):
            output.append(f'<{tag} class="{role}"{_lang_attr(role, mode, source_language)}>'
                          f'{_escape_with_refs(text, plan, block, role, ctx=ctx)}</{tag}>')
    elif kind in {"paragraph", "code"}:
        tag = "pre" if kind == "code" else "p"
        for role, text in _text_roles(block, entry, mode):
            output.append(f'<{tag} class="{role}"{_lang_attr(role, mode, source_language)}>'
                          f'{_escape_with_refs(text, plan, block, role, ctx=ctx)}</{tag}>')
    elif kind == "footnote":
        for role, text in _text_roles(block, entry, mode):
            output.append(f'<p class="footnote {role}"{_lang_attr(role, mode, source_language)}>'
                          f'{_escape_with_refs(text, plan, block, role, ctx=ctx)}</p>')
    elif kind == "list":
        for role, values in _list_roles(block, entry, mode):
            attr = _lang_attr(role, mode, source_language)
            items = "".join(
                f"<li>{_escape_with_refs(value, plan, block, role, by_offset=False, ctx=ctx)}</li>"
                for value in values
            )
            output.append(f'<ul class="{role}"{attr}>{items}</ul>')
    elif kind == "equation":
        output.append(f'<pre class="equation source">{html.escape(block.source_text)}</pre>')
    elif kind == "table":
        if block.cells:
            for role in _table_roles(block, entry, mode):
                output.append(f'<div class="table-{role}">{_table_html(block, entry, role)}</div>')
        else:
            output.append(str(block.format.get("table_html", "") or ""))
    elif kind == "image":
        for resource_id in block.resource_refs:
            filename = assets.get(resource_id)
            if filename:
                output.append(f'<img src="assets/{html.escape(filename)}" alt=""/>')
        for role, captions in _caption_roles(block, entry, mode):
            for caption in captions:
                output.append(f'<figcaption class="{role}"{_lang_attr(role, mode, source_language)}>{_html_text(caption)}</figcaption>')
    else:
        if block.source_text:
            output.append(f'<p class="source">{_html_text(block.source_text)}</p>')
    output.append("</section>")
    return "".join(output)


def _plain_superscripts(title: str) -> str:
    """章节标题（``<title>`` / 导航用）里不留 ``^{…}``，只保留上标内容本身。"""
    return SUPERSCRIPT_RE.sub(lambda match: superscript_content(match.group(1)), title or "")


def _page_html(title: str, language: str, body: str, *, stylesheet: str) -> str:
    return (
        "<!doctype html>\n"
        f'<html lang="{html.escape(language)}"><head><meta charset="utf-8"/>'
        '<meta name="viewport" content="width=device-width,initial-scale=1"/>'
        f"<title>{html.escape(_plain_superscripts(title))}</title>"
        f'<link rel="stylesheet" href="{html.escape(stylesheet)}"/></head>'
        f"<body><main>{body}</main></body></html>\n"
    )


def _xhtml(title: str, language: str, body: str) -> str:
    return (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        "<!DOCTYPE html>\n"
        '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
        f'lang="{html.escape(language)}" xml:lang="{html.escape(language)}">'
        f"<head><title>{html.escape(_plain_superscripts(title))}</title>"
        '<link rel="stylesheet" type="text/css" href="../styles/book.css"/></head>'
        f"<body>{body}</body></html>"
    )


# ---------------------------------------------------------------------------
# 各格式写入器（返回 (Path, single_file) 列表）
# ---------------------------------------------------------------------------

def _write_markdown(project: Path, root: Path, document: Document, blocks: list[Block], layer: TranslationLayer | None,
                    mode: ExportMode, assets: dict[str, str], name: str = "") -> list[tuple[Path, bool]]:
    root.mkdir(parents=True, exist_ok=True)
    # 主文件用文档名：显式 filename 优先，缺省回退 document.title（清洗后为空才退回 full）。
    target = root / f"{safe_name(name or document.title) or 'full'}.md"
    target.write_text(render_markdown(document, blocks, layer, mode, assets, build_plan(blocks)), encoding="utf-8")
    files: list[tuple[Path, bool]] = [(target, True)]
    for path in _copy_assets(project, root / "assets", document, assets):
        files.append((path, True))
    return files


def _write_html(project: Path, root: Path, document: Document, blocks: list[Block], layer: TranslationLayer | None,
                mode: ExportMode, assets: dict[str, str], name: str = "") -> list[tuple[Path, bool]]:
    root.mkdir(parents=True, exist_ok=True)
    # 主入口与其它格式一致用文档名（index.html 改为 <名称>.html）；章节/样式/assets 等
    # 包内辅助文件保持格式所需命名，章内互链不指向主入口，改后导航/锚链接仍然有效。
    entry_name = f"{safe_name(name or document.title) or 'index'}.html"
    sections = split_sections(document, blocks) or [Section(document.title, 1, blocks, nav=False)]
    language = _display_language(document, layer, mode)
    files: list[tuple[Path, bool]] = []
    nav_items: list[tuple[str, str, int]] = []
    plan = build_plan(blocks)
    ctx = _RefContext()
    source_language = str(document.source_language or "")
    section_notes = [plan.section_note_ids(section.blocks) for section in sections]
    for index, notes in enumerate(section_notes, 1):
        for note_id in notes:
            ctx.note_file[note_id] = f"chapter-{index:04d}.html"
    for index, (section, notes) in enumerate(zip(sections, section_notes), 1):
        filename = f"chapter-{index:04d}.html"
        ctx.current = filename
        body = "".join(
            _block_html(block, entry_for(layer, block), mode, assets, source_language, plan, ctx)
            for block in section.blocks
        )
        # 注文统一放章节末尾（与 EPUB 一致）。
        body += _epub_footnote_asides(notes, plan, layer, mode, source_language, ctx)
        (root / filename).write_text(_page_html(section.title, language, body, stylesheet="style.css"), encoding="utf-8")
        files.append((root / filename, False))
        if section.nav:
            nav_items.append((section.title, filename, section.level))
    if not nav_items:
        nav_items = [(section.title, f"chapter-{i:04d}.html", section.level) for i, section in enumerate(sections, 1)]
    nav_html = _render_nav(_build_nav_tree([(t, h, lv) for t, h, lv in nav_items]))
    index_body = (
        f'<h1 class="book-title">{html.escape(document.title)}</h1>'
        f'<p class="meta">{html.escape(language)} · {len(blocks)} blocks</p>'
        f'<nav class="toc">{nav_html or "<ol></ol>"}</nav>'
    )
    (root / entry_name).write_text(_page_html(document.title, language, index_body, stylesheet="style.css"), encoding="utf-8")
    files.append((root / entry_name, True))
    (root / "style.css").write_text(BASE_CSS + _INDEX_CSS_EXTRA + "\n", encoding="utf-8")
    files.append((root / "style.css", False))
    for path in _copy_assets(project, root / "assets", document, assets):
        files.append((path, False))
    return files


def _cover_jpeg(project: Path) -> tuple[str, bytes] | None:
    """用原书 PDF 首页渲染 EPUB 封面（JPEG）；非 PDF 源或渲染失败返回 None。"""
    source_dir = project / "source"
    if not source_dir.is_dir():
        return None
    pdf_path = next((child for child in sorted(source_dir.iterdir())
                     if child.is_file() and child.suffix.casefold() == ".pdf"), None)
    if pdf_path is None:
        return None
    try:
        import fitz

        with fitz.open(pdf_path) as pdf:
            if pdf.page_count < 1:
                return None
            page = pdf.load_page(0)
            zoom = max(1.0, 1600.0 / max(1.0, float(page.rect.width)))
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            if pix.width > 2400:
                scale = 2400.0 / pix.width
                pix = page.get_pixmap(matrix=fitz.Matrix(zoom * scale, zoom * scale), alpha=False)
            return "cover.jpg", pix.tobytes("jpg", jpg_quality=88)
    except Exception:  # noqa: BLE001 - 封面是增强项，任何失败都不阻断导出
        return None


def validate_epub(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if not infos or infos[0].filename != "mimetype" or infos[0].compress_type != zipfile.ZIP_STORED:
            raise StorageError("EPUB mimetype 必须是首个且未压缩条目。")
        names = set(archive.namelist())
        required = {"mimetype", "META-INF/container.xml", "EPUB/package.opf", "EPUB/nav.xhtml"}
        if missing := required - names:
            raise StorageError(f"EPUB 缺少文件：{sorted(missing)}")
        for name in names:
            if name.endswith((".xml", ".opf", ".xhtml")):
                ET.fromstring(archive.read(name))


def _write_epub(project: Path, root: Path, document: Document, blocks: list[Block], layer: TranslationLayer | None,
                mode: ExportMode, assets: dict[str, str], name: str) -> list[tuple[Path, bool]]:
    root.mkdir(parents=True, exist_ok=True)
    output = root / f"{name}.epub"
    sections = split_sections(document, blocks) or [Section(document.title, 1, blocks, nav=False)]
    language = _display_language(document, layer, mode)
    source_language = str(document.source_language or "")
    if language == source_language:
        source_language = ""  # 已由根元素声明，不必逐段重复
    manifest_items = [
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
        '<item id="css" href="styles/book.css" media-type="text/css"/>',
    ]
    cover = _cover_jpeg(project)
    if cover is not None:
        manifest_items.append('<item id="cover-image" href="cover.jpg" media-type="image/jpeg" properties="cover-image"/>')
        manifest_items.append('<item id="cover-page" href="cover.xhtml" media-type="application/xhtml+xml"/>')
    spine: list[str] = []
    if cover is not None:
        spine.append('<itemref idref="cover-page" linear="yes"/>')
    plan = build_plan(blocks)
    ctx = _RefContext()  # 同一注文可被多次引用：aside 只在首锚章节，其余跨文件链接
    section_notes = [plan.section_note_ids(section.blocks) for section in sections]
    for index, notes in enumerate(section_notes, 1):
        for note_id in notes:
            ctx.note_file[note_id] = f"text/section-{index:04d}.xhtml"
    nav_items: list[tuple[str, str, int]] = []
    bodies: list[tuple[str, str]] = []
    for index, (section, notes) in enumerate(zip(sections, section_notes), 1):
        filename = f"text/section-{index:04d}.xhtml"
        ctx.current = filename
        body = "".join(
            _block_html(block, entry_for(layer, block), mode, assets, source_language, plan, ctx)
            for block in section.blocks
        )
        # 所有注文统一放在章节末尾（章末注），不在正文流里逐段插入。
        body += _epub_footnote_asides(notes, plan, layer, mode, source_language, ctx)
        bodies.append((filename, _xhtml(section.title, language, body)))
        manifest_items.append(f'<item id="section-{index}" href="{filename}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="section-{index}"/>')
        if section.nav:
            nav_items.append((_section_nav_title(section, layer, mode), filename, section.level))
    if not nav_items:
        nav_items = [(_section_nav_title(section, layer, mode), f"text/section-{i:04d}.xhtml", section.level)
                     for i, section in enumerate(sections, 1)]
    nav_tree = _build_nav_tree(nav_items)
    nav_body = (
        '<nav xmlns:epub="http://www.idpf.org/2007/ops" epub:type="toc">'
        f"<h1>目录</h1>{_render_nav(nav_tree)}</nav>"
    )
    book_id = f"urn:uuid:{_uuid_like(document)}"
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        if cover is not None:
            cover_filename, cover_bytes = cover
            archive.writestr("EPUB/cover.jpg", cover_bytes, compress_type=zipfile.ZIP_STORED)
            archive.writestr(
                "EPUB/cover.xhtml",
                _xhtml("Cover", language, f'<div class="cover"><img src="{cover_filename}" alt="{html.escape(document.title)}"/></div>'),
            )
        archive.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
            '<rootfile full-path="EPUB/package.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>",
        )
        archive.writestr("EPUB/styles/book.css", EPUB_CSS)
        for filename, body in bodies:
            archive.writestr(f"EPUB/{filename}", body)
        for resource_id, filename in assets.items():
            relative = str((document.resources.get(resource_id) or {}).get("path") or "")
            source = project / relative if relative else None
            if source is None or not source.is_file():
                continue
            media_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
            archive.write(source, f"EPUB/text/assets/{filename}")
            manifest_items.append(
                f'<item id="asset-{len(manifest_items)}" href="text/assets/{html.escape(filename)}" media-type="{media_type}"/>'
            )
        archive.writestr("EPUB/nav.xhtml", _xhtml("目录", language, nav_body))
        package = (
            '<?xml version="1.0" encoding="utf-8"?>\n'
            f'<package xmlns="http://www.idpf.org/2007/opf" unique-identifier="book-id" version="3.0" xml:lang="{html.escape(language)}">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            f'<dc:identifier id="book-id">{book_id}</dc:identifier>'
            f"<dc:title>{html.escape(document.title)}</dc:title>"
            f"<dc:language>{html.escape(language)}</dc:language>"
            f'<meta property="dcterms:modified">{utc_now().split(".")[0]}Z</meta>'
            + (f'<meta name="cover" content="cover-image"/>' if cover is not None else "")
            + f"</metadata><manifest>{''.join(manifest_items)}</manifest>"
            f"<spine>{''.join(spine)}</spine></package>"
        )
        archive.writestr("EPUB/package.opf", package)
    validate_epub(output)
    return [(output, True)]


def _uuid_like(document: Document) -> str:
    import uuid as _uuid

    seed = f"{document.project_id}:{document.title}:{document.updated_at}"
    return str(_uuid.uuid5(_uuid.NAMESPACE_URL, seed))


def _docx_quote(paragraph: Any) -> None:
    from docx.shared import Inches, RGBColor

    paragraph.paragraph_format.left_indent = Inches(0.3)
    for run in paragraph.runs:
        run.italic = True
        run.font.color.rgb = RGBColor(0x2F, 0x51, 0x42)


def _split_superscripts(text: str) -> list[tuple[str, str]]:
    """把文本按 ``^{…}`` 切开：``("text", 片段)`` / ``("sup", 上标)``。"""
    parts: list[tuple[str, str]] = []
    cursor = 0
    for match in SUPERSCRIPT_RE.finditer(text or ""):
        if match.start() > cursor:
            parts.append(("text", text[cursor:match.start()]))
        parts.append(("sup", superscript_content(match.group(1))))
        cursor = match.end()
    if cursor < len(text or ""):
        parts.append(("text", text[cursor:]))
    return parts or [("text", text or "")]


def _docx_add_text(paragraph: Any, text: str, *, italic: bool = False) -> None:
    """写一段文本：``^{…}`` 转成上标 run（与印刷版一致），不再留 LaTeX 形态。"""
    from docx.shared import RGBColor

    for kind, value in _split_superscripts(text):
        if not value:
            continue
        run = paragraph.add_run(value)
        if kind == "sup":
            run.font.superscript = True
        if italic:
            run.italic = True
            run.font.color.rgb = RGBColor(0x2F, 0x51, 0x42)


def _docx_paragraph_with_refs(word: Any, segments: list[tuple[str, str, str]],
                              footnote_marks: dict[str, tuple[str, str]], style: Any = None) -> Any:
    """带脚注引用的段落：文本段照常，ref 段插入上标 ``w:footnoteReference``。

    用 ``w:customMarkFollows`` 保留印刷版编号（书里按章重新编号），不依赖 Word 的
    全文连续自动编号，避免正文/注文对不上。Word 不允许同一个脚注被引用两次，
    重复引用降级为普通上标（与印刷版观感一致），由调用方在 ``segments`` 里改写。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = word.add_paragraph(style=style) if style is not None else word.add_paragraph()
    for kind, value, note_id in segments:
        if kind == "text":
            _docx_add_text(paragraph, value)
        elif kind == "sup":
            run = paragraph.add_run(value)
            run.font.superscript = True
        elif note_id in footnote_marks:
            footnote_id, label = footnote_marks[note_id]
            run = paragraph.add_run()
            run.font.superscript = True
            reference = OxmlElement("w:footnoteReference")
            reference.set(qn("w:customMarkFollows"), "1")
            reference.set(qn("w:id"), footnote_id)
            run._r.append(reference)
            mark = OxmlElement("w:t")
            mark.set(qn("xml:space"), "preserve")
            mark.text = label
            run._r.append(mark)
    return paragraph


def _docx_unique_refs(segments: list[tuple[str, str, str]], seen: set[str]) -> list[tuple[str, str, str]]:
    """同一个脚注只能被 Word 引用一次；后续同号引用降级为上标文字。"""
    output: list[tuple[str, str, str]] = []
    for kind, value, note_id in segments:
        if kind != "ref" or not note_id:
            output.append((kind, value, note_id))
            continue
        if note_id in seen:
            output.append(("sup", value, ""))
        else:
            seen.add(note_id)
            output.append((kind, value, note_id))
    return output


def _docx_default_font(word: Any) -> None:
    try:
        from docx.oxml.ns import qn

        style = word.styles["Normal"]
        style.font.name = "Times New Roman"
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), "Songti SC")
    except Exception:  # noqa: BLE001 - 字体仅为观感，失败不影响内容
        pass


def _docx_table(word: Any, block: Block, entry: TranslationEntry | None, role: str) -> None:
    matrix = _table_matrix(block, entry, role)
    if not matrix:
        fallback = str(block.format.get("table_html", "") or "")
        if fallback:
            word.add_paragraph(fallback)
        return
    rows = len(matrix)
    columns = len(matrix[0])
    table = word.add_table(rows=rows, cols=columns)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    values = entry.cells if entry and entry.cells and role == "translation" else {}
    for cell in block.cells:
        row = max(0, min(rows - 1, cell.row))
        column = max(0, min(columns - 1, cell.column))
        table.cell(row, column).text = values.get(cell.cell_id, cell.text)
    for cell in block.cells:
        if cell.rowspan <= 1 and cell.colspan <= 1:
            continue
        row_end = max(0, min(rows - 1, cell.row + max(1, cell.rowspan) - 1))
        column_end = max(0, min(columns - 1, cell.column + max(1, cell.colspan) - 1))
        try:
            anchor = table.cell(max(0, min(rows - 1, cell.row)), max(0, min(columns - 1, cell.column)))
            other = table.cell(row_end, column_end)
            if anchor is not other:
                anchor.merge(other)
        except Exception:  # noqa: BLE001 - 合并失败不阻断导出
            continue


def _docx_runs_xml(text: str, *, italic: bool = False) -> str:
    """脚注部件用：文本按 ``^{…}`` 拆成普通/上标 run（``w:t`` XML 片段）。"""
    run_pr = '<w:rPr><w:i/><w:color w:val="2F5142"/></w:rPr>' if italic else ""
    parts: list[str] = []
    for kind, value in _split_superscripts(text):
        if not value:
            continue
        rpr = '<w:rPr><w:vertAlign w:val="superscript"/></w:rPr>' if kind == "sup" else run_pr
        parts.append(f'<w:r>{rpr}<w:t xml:space="preserve">{html.escape(value)}</w:t></w:r>')
    return "".join(parts)


def _docx_footnotes_part(word: Any, plan: FootnotePlan, blocks: list[Block],
                         layer: TranslationLayer | None, mode: ExportMode) -> None:
    """把已绑定页下注写入 ``word/footnotes.xml`` 并建立关系；无绑定注时不建 part。"""
    from docx.opc.constants import CONTENT_TYPE as CT
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI

    note_ids = [block.block_id for block in blocks if block.block_id in plan.links]
    if not note_ids:
        return
    entries: list[str] = [
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/>'
        '</w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>',
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
        'w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>',
    ]
    for ordinal, note_id in enumerate(note_ids, 1):
        link = plan.links[note_id]
        block = link.block
        paragraphs_xml: list[str] = []
        for role, text in _note_role_texts(link, plan, layer, mode):
            if not text:
                continue
            stripped = link.body if role == "source" else text
            if role == "source":
                paragraphs_xml.append(
                    f'<w:p><w:pPr><w:spacing w:after="60"/></w:pPr>'
                    f'<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
                    f'<w:t xml:space="preserve">{html.escape(link.number)}</w:t></w:r>'
                    f'<w:r><w:t xml:space="preserve"> </w:t></w:r>'
                    f'{_docx_runs_xml(stripped)}</w:p>'
                )
            else:
                paragraphs_xml.append(
                    f'<w:p><w:pPr><w:spacing w:after="60"/></w:pPr>'
                    f'{_docx_runs_xml(stripped, italic=True)}</w:p>'
                )
        if not paragraphs_xml:
            paragraphs_xml.append(
                f'<w:p><w:pPr><w:spacing w:after="60"/></w:pPr><w:r><w:footnoteRef/></w:r>'
                f'<w:r><w:t xml:space="preserve"> </w:t></w:r></w:p>'
            )
        entries.append(f'<w:footnote w:id="{ordinal}">{"".join(paragraphs_xml)}</w:footnote>')
    blob = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'{"".join(entries)}</w:footnotes>'
    ).encode("utf-8")
    part = Part(
        PackURI("/word/footnotes.xml"),
        CT.WML_FOOTNOTES,
        blob=blob,
        package=word.part.package,
    )
    word.part.relate_to(part, RT.FOOTNOTES)


def _write_docx(project: Path, root: Path, document: Document, blocks: list[Block], layer: TranslationLayer | None,
                mode: ExportMode, assets: dict[str, str], name: str) -> list[tuple[Path, bool]]:
    try:
        from docx import Document as WordDocument
        from docx.shared import Inches
    except ImportError as error:  # pragma: no cover - 环境缺少依赖时的明确报错
        raise StorageError("导出 DOCX 需要 python-docx。") from error
    root.mkdir(parents=True, exist_ok=True)
    output = root / f"{name}.docx"
    word = WordDocument()
    word.core_properties.title = document.title
    _docx_default_font(word)
    plan = build_plan(blocks)
    footnote_marks = {
        block.block_id: (str(ordinal), plan.links[block.block_id].number)
        for ordinal, block in enumerate((item for item in blocks if item.block_id in plan.links), 1)
    }
    seen_notes: set[str] = set()  # Word 不允许同一脚注被引用两次
    for block in blocks:
        entry = entry_for(layer, block)
        kind = block.type
        if kind == "title":
            level = max(1, min(9, int(block.level or 1)))
            style = word.styles[f"Heading {level}"]
            for role, text in _text_roles(block, entry, mode):
                if not text:
                    continue
                segments = _docx_unique_refs(annotate_footnote_markers(text, plan, block, role), seen_notes)
                heading = _docx_paragraph_with_refs(word, segments, footnote_marks, style=style)
                if mode == "bilingual" and role == "translation":
                    _docx_quote(heading)
        elif kind == "footnote":
            if block.block_id in plan.links or block.block_id in plan.absorbed:
                continue  # 已写入 word/footnotes.xml（续行块已并入前注），正文流不再重复
            for role, text in _text_roles(block, entry, mode):
                if not text:
                    continue
                segments = _docx_unique_refs(annotate_footnote_markers(text, plan, block, role), seen_notes)
                paragraph = _docx_paragraph_with_refs(word, segments, footnote_marks)
                if role == "translation" and mode != "source":
                    _docx_quote(paragraph)
        elif kind == "code":
            for role, text in _text_roles(block, entry, mode):
                if not text:
                    continue
                paragraph = word.add_paragraph(text)
                if role == "translation" and mode != "source":
                    _docx_quote(paragraph)
        elif kind == "paragraph":
            for role, text in _text_roles(block, entry, mode):
                if not text:
                    continue
                segments = _docx_unique_refs(annotate_footnote_markers(text, plan, block, role), seen_notes)
                paragraph = _docx_paragraph_with_refs(word, segments, footnote_marks)
                if role == "translation" and mode != "source":
                    _docx_quote(paragraph)
        elif kind == "list":
            for role, values in _list_roles(block, entry, mode):
                for value in values:
                    segments = _docx_unique_refs(
                        annotate_footnote_markers(value, plan, block, role, by_offset=False), seen_notes)
                    paragraph = _docx_paragraph_with_refs(word, segments, footnote_marks, style="List Bullet")
                    if mode == "bilingual" and role == "translation":
                        _docx_quote(paragraph)
        elif kind == "equation":
            word.add_paragraph(block.source_text)
        elif kind == "table":
            for role in _table_roles(block, entry, mode):
                _docx_table(word, block, entry, role)
        elif kind == "image":
            for resource_id in block.resource_refs:
                relative = str((document.resources.get(resource_id) or {}).get("path") or "")
                source = project / relative if relative else None
                if source is not None and source.is_file():
                    try:
                        word.add_picture(str(source), width=Inches(5.8))
                    except Exception:  # noqa: BLE001 - 个别图片格式失败时留占位说明
                        word.add_paragraph(f"[图片：{source.name}]")
                else:
                    filename = assets.get(resource_id) or resource_id
                    word.add_paragraph(f"[图片：{filename}]")
            for _role, captions in _caption_roles(block, entry, mode):
                for caption in captions:
                    paragraph = word.add_paragraph(style="Caption")
                    _docx_add_text(paragraph, caption)
        else:
            if block.source_text:
                paragraph = word.add_paragraph()
                _docx_add_text(paragraph, block.source_text)
    _docx_footnotes_part(word, plan, blocks, layer, mode)
    word.save(output)
    return [(output, True)]


def _display_language(document: Document, layer: TranslationLayer | None, mode: ExportMode) -> str:
    if mode == "source":
        return document.source_language or "und"
    if layer is not None and layer.locale:
        return layer.locale
    return document.target_language or document.source_language or "und"


# ---------------------------------------------------------------------------
# 对外 API
# ---------------------------------------------------------------------------

def available_locales(project: Path) -> list[str]:
    root = project / "translations"
    values: list[str] = []
    if not root.is_dir():
        return values
    for child in sorted(root.iterdir()):
        if not child.is_dir() or not (child / "translation.json").is_file():
            continue
        try:
            if safe_locale(child.name) == child.name:
                values.append(child.name)
        except StorageError:
            continue
    return values


def export_options(store: ProjectStore, project_id: str, locale: str = "") -> dict[str, Any]:
    project = store.project_path(project_id)
    document = store.load_document(project_id)
    locales = available_locales(project)
    requested = str(locale or "").strip()
    if requested in locales:
        default_locale = requested
    elif document.target_language in locales:
        default_locale = document.target_language
    elif locales:
        default_locale = locales[0]
    else:
        default_locale = requested or document.target_language or ""
    layer = store.load_translation(project_id, default_locale) if default_locale else None

    blocks = document.blocks
    positions = _block_positions(blocks)
    by_level = _title_positions(blocks)
    levels = toc_levels(document)
    prefix = [0]
    for block in blocks:
        prefix.append(prefix[-1] + (1 if has_translation(block, entry_for(layer, block)) else 0))

    def node_json(node: Any) -> dict[str, Any]:
        block_id = str(node.block_id)
        start = positions.get(block_id)
        if start is None:
            count = translated = 0
        else:
            level = levels.get(block_id) or 1
            end = _chapter_end(by_level, start, level, len(blocks))
            count = end - start
            translated = prefix[end] - prefix[start]
        return {
            "node_id": str(node.node_id),
            "title": str(node.title or ""),
            "level": max(1, min(6, int(node.level or 1))),
            "block_id": block_id,
            "page_index": int(getattr(node, "page_index", 0) or 0),
            "block_count": count,
            "translated_count": translated,
            "missing_count": max(0, count - translated),
            "children": [node_json(child) for child in (getattr(node, "children", []) or [])],
        }

    return {
        "project_id": project_id,
        "title": document.title,
        "source_language": document.source_language,
        "target_language": document.target_language,
        "locales": locales,
        "default_locale": default_locale,
        "formats": FORMATS,
        "modes": MODES,
        "block_count": len(blocks),
        "image_count": len(document.resources),
        "toc": [node_json(node) for node in document.toc],
        "recent": list_exports(store, project_id),
    }


def _unique_export_id(project: Path, mode: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    base = f"{stamp}-{mode}"
    candidate = base
    index = 1
    while (project / "exports" / candidate).exists():
        index += 1
        candidate = f"{base}-{index}"
    return candidate


def _format_files(
    fmt: str,
    project: Path,
    root: Path,
    document: Document,
    blocks: list[Block],
    layer: TranslationLayer | None,
    mode: ExportMode,
    assets: dict[str, str],
    name: str,
) -> list[tuple[Path, bool]]:
    if fmt == "epub":
        return _write_epub(project, root, document, blocks, layer, mode, assets, name)
    if fmt == "docx":
        return _write_docx(project, root, document, blocks, layer, mode, assets, name)
    if fmt == "markdown":
        return _write_markdown(project, root, document, blocks, layer, mode, assets, name)
    if fmt == "html":
        return _write_html(project, root, document, blocks, layer, mode, assets, name)
    raise StorageError(f"不支持的导出格式：{fmt}")


def create_export(
    store: ProjectStore,
    project_id: str,
    *,
    formats: Iterable[str] | None = None,
    mode: str = "translation",
    locale: str = "zh-CN",
    chapters: Iterable[str] | None = None,
    block_ids: Iterable[str] | None = None,
    filename: str = "",
    missing_policy: str = "fallback_source",
) -> dict[str, Any]:
    project = store.project_path(project_id)
    document = store.load_document(project_id)

    requested = list(dict.fromkeys(str(value) for value in (formats or [])))
    if not requested:
        requested = ["epub"]
    unknown_formats = [value for value in requested if value not in FORMAT_IDS]
    if unknown_formats:
        raise StorageError(f"不支持的导出格式：{'、'.join(unknown_formats)}")
    if mode not in MODE_IDS:
        raise StorageError(f"不支持的导出模式：{mode}")
    if missing_policy != "fallback_source":
        raise StorageError(f"不支持的缺失译文策略：{missing_policy}")

    layer = None if mode == "source" else store.load_translation(project_id, locale)
    chapter_values = [str(value) for value in (chapters or [])]
    block_values = [str(value) for value in (block_ids or [])]
    known = {block.block_id for block in document.blocks}
    warnings: list[str] = []
    if block_values:
        ignored = [value for value in block_values if value not in known]
        if ignored:
            warnings.append(f"忽略 {len(ignored)} 个不存在的段落。")
    elif chapter_values:
        ignored = [value for value in chapter_values if value not in known]
        if ignored:
            warnings.append(f"忽略 {len(ignored)} 个不存在的章节。")

    blocks = select_blocks(document, chapter_values, block_values)
    if not blocks:
        warnings.append("所选章节无任何可导出内容。")
    if mode != "source" and blocks:
        missing = sum(1 for block in blocks if not has_translation(block, entry_for(layer, block)))
        if missing:
            warnings.append(f"译文缺失 {missing} 段，已回退原文。")

    export_id = _unique_export_id(project, mode)
    export_dir = project / "exports" / export_id
    export_dir.mkdir(parents=True, exist_ok=True)
    name = safe_name(str(filename or "").strip() or document.title)
    assets = asset_plan(document, project)
    flat = len(requested) == 1
    written: list[tuple[Path, bool]] = []
    for fmt in requested:
        root = export_dir if flat else export_dir / fmt
        written.extend(_format_files(fmt, project, root, document, blocks, layer, mode, assets, name))

    files: list[dict[str, Any]] = []
    seen: set[str] = set()
    for path, single in written:
        try:
            relative = path.relative_to(export_dir).as_posix()
        except ValueError:
            continue
        if relative in seen:
            continue
        seen.add(relative)
        files.append({
            "name": path.name,
            "rel_path": relative,
            "bytes": path.stat().st_size if path.is_file() else 0,
            "single_file": bool(single),
        })
    files.sort(key=lambda item: item["rel_path"])
    total_bytes = sum(int(item["bytes"]) for item in files)
    manifest = {
        "export_id": export_id,
        "created_at": utc_now(),
        "mode": mode,
        "locale": "" if mode == "source" else locale,
        "formats": requested,
        "path": str(export_dir),
        "files": files,
        "bytes": total_bytes,
        "warnings": warnings,
        "title": document.title,
        "chapters": chapter_values,
        "block_ids": block_values,
        "missing_policy": missing_policy,
        "block_count": len(blocks),
    }
    atomic_json(export_dir / "export.json", manifest)
    return manifest


def _recent_projection(value: dict[str, Any]) -> dict[str, Any]:
    return {
        "export_id": str(value.get("export_id", "")),
        "created_at": str(value.get("created_at", "")),
        "display_name": str(value.get("display_name", "")),
        "mode": str(value.get("mode", "")),
        "locale": str(value.get("locale", "")),
        "formats": list(value.get("formats") or []),
        "path": str(value.get("path", "")),
        "files": list(value.get("files") or []),
        "bytes": int(value.get("bytes") or 0),
        "warnings": list(value.get("warnings") or []),
    }


def list_exports(store: ProjectStore, project_id: str) -> list[dict[str, Any]]:
    project = store.project_path(project_id)
    root = project / "exports"
    values: list[dict[str, Any]] = []
    if root.is_dir():
        for manifest_path in root.glob("*/export.json"):
            try:
                value = json.loads(manifest_path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                continue
            if isinstance(value, dict):
                values.append(_recent_projection(value))
    values.sort(key=lambda item: str(item.get("created_at", "")), reverse=True)
    return values[:50]


def _export_record(store: ProjectStore, project_id: str, export_id: str) -> tuple[Path, dict[str, Any]]:
    """校验 ``export_id`` 并返回 ``(导出目录, manifest)``；非法抛 StorageError，缺失抛 FileNotFoundError。"""
    if not _EXPORT_ID_RE.match(str(export_id or "")):
        raise StorageError("非法的导出标识。")
    project = store.project_path(project_id)
    exports_root = (project / "exports").resolve()
    base = (exports_root / export_id).resolve()
    if base == exports_root or exports_root not in base.parents:
        raise StorageError("非法的导出标识。")
    manifest_path = base / "export.json"
    if not manifest_path.is_file():
        raise FileNotFoundError(export_id)
    try:
        value = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as error:
        raise StorageError("导出记录已损坏，无法读取 export.json。") from error
    if not isinstance(value, dict):
        raise StorageError("导出记录已损坏，无法读取 export.json。")
    return base, value


def delete_export(store: ProjectStore, project_id: str, export_id: str) -> dict[str, Any]:
    """删除一次导出（含产物目录与历史记录），不影响 raw / document / translations。"""
    base, _value = _export_record(store, project_id, export_id)
    shutil.rmtree(base)
    return {"export_id": export_id, "deleted": True}


def rename_export(store: ProjectStore, project_id: str, export_id: str, display_name: str) -> dict[str, Any]:
    """重命名一次导出：更新 manifest 的 display_name，并把顶层单文件产物改成同名文件。

    文件夹类产物（assets/、分章 HTML）保持原样；名称冲突时自动加 " (n)" 后缀。
    """
    base, value = _export_record(store, project_id, export_id)
    raw = str(display_name or "").strip()
    if not raw:
        raise StorageError("导出名称不能为空。")
    name = safe_name(raw)
    files = value.get("files")
    if isinstance(files, list):
        updated: list[Any] = []
        for item in files:
            entry = dict(item) if isinstance(item, dict) else item
            if isinstance(entry, dict) and entry.get("single_file"):
                old_rel = str(entry.get("rel_path", ""))
                try:
                    old_path = (base / old_rel).resolve()
                except (OSError, ValueError):
                    old_path = None
                if old_path is not None and old_path.is_file() and base in old_path.parents:
                    parent = old_path.parent
                    candidate = parent / f"{name}{old_path.suffix}"
                    index = 2
                    while candidate.exists() and candidate != old_path:
                        candidate = parent / f"{name} ({index}){old_path.suffix}"
                        index += 1
                    if candidate != old_path:
                        old_path.rename(candidate)
                    entry["name"] = candidate.name
                    entry["rel_path"] = candidate.relative_to(base).as_posix()
            updated.append(entry)
        value["files"] = updated
    value["display_name"] = name
    atomic_json(base / "export.json", value)
    return _recent_projection(value)


def resolve_export_file(store: ProjectStore, project_id: str, export_id: str, name: str) -> tuple[Path, str]:
    """把 ``(export_id, name)`` 解析成导出目录内的真实文件，并拒绝目录穿越。

    非法输入抛 ``StorageError``（API 层 400），文件不存在抛 ``FileNotFoundError``（API 层 404）。
    """
    if not _EXPORT_ID_RE.match(str(export_id or "")):
        raise StorageError("非法的导出标识。")
    if not name or "\x00" in name:
        raise StorageError("非法的文件名。")
    project = store.project_path(project_id)
    exports_root = (project / "exports").resolve()
    base = (exports_root / export_id).resolve()
    if base == exports_root or exports_root not in base.parents:
        raise StorageError("非法的导出标识。")
    candidate = (base / name).resolve()
    if candidate != base and base not in candidate.parents:
        raise StorageError("文件超出导出目录范围。")
    if candidate == base or not candidate.is_file():
        raise FileNotFoundError(name)
    extension = candidate.suffix.casefold()
    media_type = MEDIA_TYPES.get(extension) or mimetypes.guess_type(candidate.name)[0] or "application/octet-stream"
    return candidate, media_type
