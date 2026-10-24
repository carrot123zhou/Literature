"""内置多格式导出：渲染层单测 + API 契约测试。

覆盖：三模式 × 四格式、章节子集展开（含嵌套层级边界）、缺译回退与不重复、
CJK/emoji 往返、表格 cells + rowspan/colspan、图片 basename 冲突去重、
EPUB 结构校验、DOCX 可打开且含 document.xml、空章节/空文档、export_v1 能力标记、
路径穿越被拒、TBL 四个端点仍为 410。
"""

from __future__ import annotations

import json
import re
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from literature.exporters import (
    FORMAT_IDS,
    asset_plan,
    create_export,
    export_options,
    render_markdown,
    select_blocks,
    split_sections,
    validate_epub,
)
from literature.footnote_refs import build_plan
from literature.models import Block, Cell, Document, TocNode, TranslationEntry, TranslationLayer
from literature.storage import ProjectStore


CJK = "中文标题😀"
EMOJI_PARAGRAPH = "Emoji 😀 与 CJK 中文混排。"


def _blocks() -> list[Block]:
    return [
        Block("t1", "title", 0, 0, source_text="Chapter One", level=1),
        Block("t1a", "title", 1, 0, source_text="Section One A", level=2),
        Block("p1", "paragraph", 2, 0, source_text="Hello world."),
        Block("p2", "paragraph", 3, 0, source_text="No translation here."),
        Block("l1", "list", 4, 0, items=["alpha", "beta"]),
        Block("tb", "table", 5, 0, cells=[
            Cell("c1", 0, 0, "A", header=True),
            Cell("c2", 0, 1, "B"),
            Cell("c3", 1, 0, "C", colspan=2),
        ]),
        Block("img", "image", 6, 0, resource_refs=["r1", "r2", "missing"], caption=["Figure 1"]),
        Block("eq", "equation", 7, 0, source_text="E=mc^2"),
        Block("fn", "footnote", 8, 0, source_text="A note."),
        Block("t2", "title", 9, 0, source_text="Chapter Two", level=1),
        Block("p3", "paragraph", 10, 0, source_text="Second chapter."),
        Block("cjk", "paragraph", 11, 0, source_text=EMOJI_PARAGRAPH),
    ]


def _document(project_root: Path) -> Document:
    assets = project_root / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    (assets / "pic.png").write_bytes(b"\x89PNG\r\n\x1a\n1")
    (assets / "other").mkdir(parents=True, exist_ok=True)
    (assets / "other" / "pic.png").write_bytes(b"\x89PNG\r\n\x1a\n2")
    return Document(
        project_id="placeholder",
        title="Book",
        source_language="en",
        target_language="zh-CN",
        blocks=_blocks(),
        toc=[
            TocNode("n1", "Chapter One", 1, "t1", 0, [TocNode("n1a", "Section One A", 2, "t1a", 0, [])]),
            TocNode("n2", "Chapter Two", 1, "t2", 0, []),
        ],
        resources={
            "r1": {"path": "assets/pic.png"},
            "r2": {"path": "assets/other/pic.png"},
            "missing": {"path": "assets/none.png"},
        },
    )


def _layer(project_id: str) -> TranslationLayer:
    return TranslationLayer(
        project_id=project_id,
        locale="zh-CN",
        entries={
            "t1": TranslationEntry("t1", text="第一章"),
            "t1a": TranslationEntry("t1a", text="第一节"),
            "p1": TranslationEntry("p1", text="你好，世界。"),
            "l1": TranslationEntry("l1", items=["甲", "乙"]),
            "tb": TranslationEntry("tb", cells={"c1": "甲", "c2": "乙", "c3": "丙"}),
            "img": TranslationEntry("img", caption=["图 1"]),
            "t2": TranslationEntry("t2", text="第二章"),
            "cjk": TranslationEntry("cjk", text="译文：" + EMOJI_PARAGRAPH),
        },
    )


class SelectionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_chapter_expansion_includes_nested_and_stops_at_next_title(self) -> None:
        document = _document(self.root)
        ids = [b.block_id for b in select_blocks(document, chapters=["t1"])]
        self.assertEqual(ids, ["t1", "t1a", "p1", "p2", "l1", "tb", "img", "eq", "fn"])
        nested = [b.block_id for b in select_blocks(document, chapters=["t1a"])]
        self.assertEqual(nested, ["t1a", "p1", "p2", "l1", "tb", "img", "eq", "fn"])
        second = [b.block_id for b in select_blocks(document, chapters=["t2"])]
        self.assertEqual(second, ["t2", "p3", "cjk"])

    def test_chapter_union_merges_ranges(self) -> None:
        document = _document(self.root)
        ids = [b.block_id for b in select_blocks(document, chapters=["t1", "t2"])]
        self.assertEqual(ids, [b.block_id for b in document.blocks])

    def test_empty_chapters_means_whole_book(self) -> None:
        document = _document(self.root)
        self.assertEqual(len(select_blocks(document)), len(document.blocks))
        self.assertEqual(len(select_blocks(document, chapters=[])), len(document.blocks))

    def test_block_ids_override_keeps_reading_order(self) -> None:
        document = _document(self.root)
        ids = [b.block_id for b in select_blocks(document, block_ids=["p3", "p1", "p1"])]
        self.assertEqual(ids, ["p1", "p3"])

    def test_split_sections_uses_toc_hierarchy(self) -> None:
        document = _document(self.root)
        sections = split_sections(document, document.blocks)
        self.assertEqual([(s.title, s.level, s.nav) for s in sections], [
            ("Chapter One", 1, True),
            ("Section One A", 2, True),
            ("Chapter Two", 1, True),
        ])
        self.assertEqual([b.block_id for b in sections[0].blocks], ["t1"])
        self.assertEqual([b.block_id for b in sections[1].blocks], ["t1a", "p1", "p2", "l1", "tb", "img", "eq", "fn"])

    def test_empty_document_sections(self) -> None:
        document = _document(self.root)
        document.blocks = []
        document.toc = []
        self.assertEqual(split_sections(document, []), [])


class RenderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.document = _document(self.root)
        self.layer = _layer("p")
        self.assets = asset_plan(self.document, self.root)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_source_mode_only_emits_source(self) -> None:
        text = render_markdown(self.document, self.document.blocks, self.layer, "source", self.assets)
        self.assertIn("# Chapter One", text)
        self.assertIn("Hello world.", text)
        self.assertNotIn("你好，世界。", text)

    def test_translation_mode_falls_back_to_source(self) -> None:
        text = render_markdown(self.document, self.document.blocks, self.layer, "translation", self.assets)
        self.assertIn("你好，世界。", text)
        self.assertIn("No translation here.", text)
        self.assertNotIn("Hello world.", text)

    def test_bilingual_interleaves_without_duplicate_fallback(self) -> None:
        text = render_markdown(self.document, self.document.blocks, self.layer, "bilingual", self.assets)
        self.assertIn("Hello world.", text)
        self.assertIn("> 你好，世界。", text)
        self.assertIn("No translation here.", text)
        self.assertNotIn("> No translation here.", text)
        self.assertIn("- alpha", text)
        self.assertIn("> - 甲", text)

    def test_table_cells_and_spans(self) -> None:
        text = render_markdown(self.document, self.document.blocks, self.layer, "translation", self.assets)
        self.assertIn("甲", text)
        self.assertIn("丙", text)
        from literature.exporters import _table_html

        table = self.document.block("tb")
        entry = self.layer.entries["tb"]
        markup = _table_html(table, entry, "translation")
        self.assertIn('colspan="2"', markup)
        self.assertIn("<th>甲</th>", markup)
        self.assertIn("<td colspan=\"2\">丙</td>", markup)

    def test_asset_basename_conflict_is_deduplicated(self) -> None:
        names = sorted(self.assets.values())
        self.assertEqual(len(names), 2)
        self.assertEqual(len(set(names)), 2)
        self.assertIn("pic.png", names)
        text = render_markdown(self.document, self.document.blocks, self.layer, "source", self.assets)
        for name in names:
            self.assertIn(f"assets/{name}", text)
        self.assertNotIn("assets/none.png", text)

    def test_cjk_and_emoji_roundtrip_in_markdown(self) -> None:
        text = render_markdown(self.document, self.document.blocks, self.layer, "translation", self.assets)
        self.assertIn(EMOJI_PARAGRAPH, text)
        self.assertIn("译文：" + EMOJI_PARAGRAPH, text)


class ApiExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        source = self.root / "source.html"
        source.write_text("<h1>Book</h1>", encoding="utf-8")
        self.store = ProjectStore(self.root / "library")
        created = self.store.create_project(source, title="Book")
        self.project_id = created["project_id"]
        self.project_dir = Path(created["path"])
        self.document = _document(self.project_dir)
        self.document.project_id = self.project_id
        self.store.save_document(self.document)
        self.store.save_translation(_layer(self.project_id))

        from fastapi.testclient import TestClient

        from literature.api import create_app

        self.client = TestClient(create_app(self.store))
        self.headers = {"x-literature-token": self.client.app.state.session_token}

    def tearDown(self) -> None:
        self.client.close()
        self.temporary.cleanup()

    def _export(self, **payload):
        return self.client.post(f"/projects/{self.project_id}/export", headers=self.headers, json=payload)

    def test_capability_marker_and_tbl_still_gone(self) -> None:
        health = self.client.get("/health").json()
        self.assertIn("export_v1", health["capabilities"])
        for method, path in (
            ("post", f"/projects/{self.project_id}/tbl/prepare"), ("post", "/tbl/start"),
            ("get", "/tbl/config"), ("put", "/tbl/config"), ("get", "/tbl/status"),
        ):
            response = getattr(self.client, method)(path, headers=self.headers)
            self.assertEqual(response.status_code, 410, response.text)

    def test_options_reports_locales_formats_modes_and_counts(self) -> None:
        response = self.client.get(f"/projects/{self.project_id}/export/options", headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        value = response.json()
        self.assertEqual({item["id"] for item in value["formats"]}, FORMAT_IDS)
        self.assertEqual({item["id"] for item in value["modes"]}, {"source", "translation", "bilingual"})
        self.assertEqual(value["locales"], ["zh-CN"])
        self.assertEqual(value["default_locale"], "zh-CN")
        self.assertEqual(value["block_count"], len(self.document.blocks))
        first = value["toc"][0]
        self.assertEqual(first["title"], "Chapter One")
        self.assertEqual(first["block_count"], 9)
        self.assertEqual(first["translated_count"], 6)
        self.assertEqual(first["missing_count"], 3)
        self.assertEqual(first["children"][0]["title"], "Section One A")

    def test_all_formats_and_modes_produce_files(self) -> None:
        for mode in ("source", "translation", "bilingual"):
            response = self._export(formats=["epub", "docx", "markdown", "html"], mode=mode, locale="zh-CN")
            self.assertEqual(response.status_code, 200, response.text)
            value = response.json()
            rels = {item["rel_path"] for item in value["files"]}
            self.assertIn("epub/Book.epub", rels)
            self.assertIn("docx/Book.docx", rels)
            self.assertIn("markdown/Book.md", rels)
            self.assertIn("html/Book.html", rels)
            self.assertTrue(any(rel.startswith("html/chapter-") for rel in rels))
            # 主入口与其它格式一致：主文件为 single_file（可另存/可重命名）。
            single = {item["rel_path"] for item in value["files"] if item["single_file"]}
            self.assertLessEqual({"epub/Book.epub", "docx/Book.docx", "markdown/Book.md", "html/Book.html"}, single)
            self.assertFalse(any(rel in single for rel in rels if rel.startswith("html/chapter-") or rel.endswith("style.css")))
            for item in value["files"]:
                self.assertTrue((Path(value["path"]) / item["rel_path"]).is_file())
            markdown = (Path(value["path"]) / "markdown" / "Book.md").read_text(encoding="utf-8")
            if mode == "source":
                self.assertIn("Hello world.", markdown)
                self.assertNotIn("你好，世界。", markdown)
            else:
                self.assertIn("你好，世界。", markdown)
                self.assertIn("No translation here.", markdown)

    def test_single_format_is_flat_and_uses_filename(self) -> None:
        response = self._export(formats=["epub"], mode="source", filename="custom name")
        self.assertEqual(response.status_code, 200, response.text)
        value = response.json()
        self.assertEqual([item["rel_path"] for item in value["files"]], ["custom name.epub"])
        self.assertEqual(value["files"][0]["single_file"], True)

    def test_chapter_subset_excludes_other_chapters(self) -> None:
        response = self._export(formats=["markdown"], mode="source", chapters=["t2"])
        self.assertEqual(response.status_code, 200, response.text)
        path = Path(response.json()["path"]) / "Book.md"
        markdown = path.read_text(encoding="utf-8")
        self.assertIn("Second chapter.", markdown)
        self.assertNotIn("Hello world.", markdown)

    def test_block_ids_override_chapters(self) -> None:
        response = self._export(formats=["markdown"], mode="source", chapters=["t1"], block_ids=["p3"])
        self.assertEqual(response.status_code, 200, response.text)
        markdown = (Path(response.json()["path"]) / "Book.md").read_text(encoding="utf-8")
        self.assertIn("Second chapter.", markdown)
        self.assertNotIn("Hello world.", markdown)

    def test_unknown_chapter_yields_warning_and_empty_selection_still_exports(self) -> None:
        response = self._export(formats=["markdown"], mode="source", chapters=["t2", "nope"])
        self.assertEqual(response.status_code, 200, response.text)
        # 选到 t2 有效内容 + 一个不存在章节 → 有告警但成功。
        self.assertTrue(any("不存在" in item for item in response.json()["warnings"]))
        empty = self._export(formats=["markdown"], mode="source", chapters=["nope"])
        self.assertEqual(empty.status_code, 200, empty.text)
        self.assertTrue(any("无任何可导出内容" in item for item in empty.json()["warnings"]))
        self.assertEqual(empty.json()["block_count"], 0)

    def test_missing_translation_warning(self) -> None:
        response = self._export(formats=["markdown"], mode="translation")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(any("译文缺失" in item for item in response.json()["warnings"]))

    def test_docx_is_openable_zip_and_contains_document_xml(self) -> None:
        response = self._export(formats=["docx"], mode="translation")
        self.assertEqual(response.status_code, 200, response.text)
        path = Path(response.json()["path"]) / "Book.docx"
        self.assertTrue(zipfile.is_zipfile(path))
        with zipfile.ZipFile(path) as archive:
            self.assertIn("word/document.xml", archive.namelist())
        from docx import Document as WordDocument

        word = WordDocument(str(path))
        text = "\n".join(paragraph.text for paragraph in word.paragraphs)
        self.assertIn("中文", text)
        self.assertIn("😀", text)
        self.assertTrue(word.tables)
        cell_texts = {cell.text for table in word.tables for row in table.rows for cell in row.cells}
        self.assertIn("甲", cell_texts)

    def test_epub_structure_is_valid(self) -> None:
        response = self._export(formats=["epub"], mode="bilingual")
        self.assertEqual(response.status_code, 200, response.text)
        path = Path(response.json()["path"]) / "Book.epub"
        validate_epub(path)
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            self.assertEqual(infos[0].filename, "mimetype")
            self.assertEqual(infos[0].compress_type, zipfile.ZIP_STORED)
            for name in ("META-INF/container.xml", "EPUB/package.opf", "EPUB/nav.xhtml"):
                self.assertIn(name, archive.namelist())
                ET.fromstring(archive.read(name))
            self.assertTrue(any(name.endswith(".xhtml") and "section" in name for name in archive.namelist()))
            css = archive.read("EPUB/styles/book.css").decode("utf-8")
            self.assertIn("text-align:justify", css)
            self.assertNotIn("border-left", css)

    def test_epub_nav_bilingual_and_translation_modes(self) -> None:
        # 对照版：导航目录为「原文 · 译文」。
        response = self._export(formats=["epub"], mode="bilingual")
        path = Path(response.json()["path"]) / "Book.epub"
        with zipfile.ZipFile(path) as archive:
            nav = archive.read("EPUB/nav.xhtml").decode("utf-8")
        self.assertIn("Chapter One · 第一章", nav)
        self.assertIn("Section One A · 第一节", nav)

        # 译文版：导航目录用纯译文标题；缺译回退原文。
        response = self._export(formats=["epub"], mode="translation")
        path = Path(response.json()["path"]) / "Book.epub"
        with zipfile.ZipFile(path) as archive:
            nav = archive.read("EPUB/nav.xhtml").decode("utf-8")
        self.assertIn("第一章", nav)
        self.assertIn("第二章", nav)
        self.assertNotIn("Chapter One", nav)

    def test_epub_cover_from_source_pdf(self) -> None:
        import fitz

        pdf = self.project_dir / "source" / "origin.pdf"
        with fitz.open() as doc:
            page = doc.new_page(width=400, height=600)
            page.insert_text((40, 80), "COVER", fontsize=32)
            # 填充一些像素，避免空白页被压缩成极小文件。
            page.draw_rect(fitz.Rect(20, 20, 380, 580), color=(0.2, 0.3, 0.2), fill=None)
            doc.save(pdf)
        response = self._export(formats=["epub"], mode="bilingual")
        self.assertEqual(response.status_code, 200, response.text)
        path = Path(response.json()["path"]) / "Book.epub"
        validate_epub(path)
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            self.assertIn("EPUB/cover.jpg", names)
            self.assertIn("EPUB/cover.xhtml", names)
            self.assertGreater(archive.getinfo("EPUB/cover.jpg").file_size, 1000)
            opf = archive.read("EPUB/package.opf").decode("utf-8")
            self.assertIn('properties="cover-image"', opf)
            self.assertIn('<meta name="cover" content="cover-image"/>', opf)
            self.assertIn('<itemref idref="cover-page"', opf)
            spine_first = opf.split("<spine>", 1)[1].split("</spine>", 1)[0]
            self.assertTrue(spine_first.index('idref="cover-page"') < spine_first.index("section-1"))

    def test_epub_without_pdf_source_has_no_cover(self) -> None:
        response = self._export(formats=["epub"], mode="bilingual")
        path = Path(response.json()["path"]) / "Book.epub"
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
        self.assertNotIn("EPUB/cover.jpg", names)
        self.assertNotIn("EPUB/cover.xhtml", names)
        with zipfile.ZipFile(path) as archive:
            opf = archive.read("EPUB/package.opf").decode("utf-8")
        self.assertNotIn("cover-image", opf)

    def test_html_is_multipage_with_nav_and_css(self) -> None:
        response = self._export(formats=["html"], mode="source")
        self.assertEqual(response.status_code, 200, response.text)
        root = Path(response.json()["path"])
        self.assertFalse((root / "index.html").exists())
        entry = (root / "Book.html").read_text(encoding="utf-8")
        self.assertIn("Chapter One", entry)
        self.assertIn("chapter-0001.html", entry)
        # 主入口里的章内链接全部存在（改名为文档名后导航仍有效）。
        for href in re.findall(r'href="([^"]+)"', entry):
            if href.startswith("http") or href.startswith("#"):  # 外链与锚点不落盘
                continue
            self.assertTrue((root / href).is_file(), href)
        self.assertTrue((root / "style.css").is_file())
        # 章节页互不指向主入口固定名（index.html），改主入口名不破坏章节内链。
        chapter = (root / "chapter-0001.html").read_text(encoding="utf-8")
        self.assertNotIn("index.html", chapter)

    def test_default_names_use_document_title_and_clean_illegal_chars(self) -> None:
        # 中文文档名直接作为主文件名；非法路径字符被 safe_name 清洗；不出现双扩展名。
        from literature.storage import safe_name
        document = self.document
        document.title = "中国的道教：研究&笔记?"
        self.store.save_document(document)
        response = self._export(formats=["markdown", "html"], mode="source")
        self.assertEqual(response.status_code, 200, response.text)
        value = response.json()
        expected = safe_name(document.title)
        rels = [item["rel_path"] for item in value["files"]]
        self.assertIn(f"markdown/{expected}.md", rels)
        self.assertIn(f"html/{expected}.html", rels)
        for rel in rels:
            self.assertFalse(rel.endswith(".md.md") or rel.endswith(".html.html"))

    def test_markdown_private_entry_falls_back_to_document_title(self) -> None:
        # 私有缺省入口（不传 name）也应回退 document.title，而非 full。
        from literature.exporters import _write_markdown
        written = _write_markdown(self.project_dir, self.root / "md-out", self.document, self.document.blocks, None, "source", {})
        self.assertEqual(written[0][0].name, "Book.md")

    def test_download_and_path_traversal(self) -> None:
        response = self._export(formats=["epub"], mode="source")
        value = response.json()
        export_id = value["export_id"]
        rel = value["files"][0]["rel_path"]
        ok = self.client.get(
            f"/projects/{self.project_id}/exports/{export_id}/file", headers=self.headers, params={"name": rel},
        )
        self.assertEqual(ok.status_code, 200, ok.text)
        self.assertEqual(ok.headers["content-type"].split(";")[0], "application/epub+zip")
        self.assertGreater(len(ok.content), 0)
        for bad in ("../../project.json", "../export.json", "/etc/passwd"):
            blocked = self.client.get(
                f"/projects/{self.project_id}/exports/{export_id}/file", headers=self.headers, params={"name": bad},
            )
            self.assertIn(blocked.status_code, (400, 404), f"{bad} -> {blocked.status_code}")
        missing = self.client.get(
            f"/projects/{self.project_id}/exports/{export_id}/file", headers=self.headers, params={"name": "nope.txt"},
        )
        self.assertEqual(missing.status_code, 404, missing.text)

    def test_exports_list_is_newest_first_and_export_does_not_mutate_sources(self) -> None:
        first = self._export(formats=["markdown"], mode="source")
        second = self._export(formats=["markdown"], mode="translation")
        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(second.status_code, 200, second.text)
        listing = self.client.get(f"/projects/{self.project_id}/exports", headers=self.headers)
        self.assertEqual(listing.status_code, 200, listing.text)
        values = listing.json()
        self.assertEqual(len(values), 2)
        self.assertEqual(values[0]["export_id"], second.json()["export_id"])
        # 导出是视图：document / translation 内容不变。
        self.assertEqual(self.store.load_document(self.project_id).blocks[2].source_text, "Hello world.")
        self.assertEqual(self.store.load_translation(self.project_id, "zh-CN").entries["p1"].text, "你好，世界。")

    def test_export_options_requires_document(self) -> None:
        source = self.root / "empty.html"
        source.write_text("<h1>x</h1>", encoding="utf-8")
        created = self.store.create_project(source, title="No Doc")
        response = self.client.get(f"/projects/{created['project_id']}/export/options", headers=self.headers)
        self.assertEqual(response.status_code, 400, response.text)


class ExportManageTests(unittest.TestCase):
    """导出管理：删除与重命名只动 exports/ 目录，不影响 raw / document / translations。"""

    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        source = self.root / "source.html"
        source.write_text("<h1>Book</h1>", encoding="utf-8")
        self.store = ProjectStore(self.root / "library")
        created = self.store.create_project(source, title="Book")
        self.project_id = created["project_id"]
        self.project_dir = Path(created["path"])
        self.document = _document(self.project_dir)
        self.document.project_id = self.project_id
        self.store.save_document(self.document)

        from fastapi.testclient import TestClient

        from literature.api import create_app

        self.client = TestClient(create_app(self.store))
        self.headers = {"x-literature-token": self.client.app.state.session_token}

    def tearDown(self) -> None:
        self.client.close()
        self.temporary.cleanup()

    def _export(self, **payload):
        return self.client.post(f"/projects/{self.project_id}/export", headers=self.headers, json=payload)

    def test_delete_removes_dir_and_history_and_keeps_sources(self) -> None:
        created = self._export(formats=["epub"], mode="bilingual")
        self.assertEqual(created.status_code, 200, created.text)
        export_id = created.json()["export_id"]
        exports_dir = self.project_dir / "exports" / export_id
        self.assertTrue(exports_dir.is_dir())

        response = self.client.delete(f"/projects/{self.project_id}/exports/{export_id}", headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), {"export_id": export_id, "deleted": True})
        self.assertFalse(exports_dir.exists())
        listing = self.client.get(f"/projects/{self.project_id}/exports", headers=self.headers)
        self.assertEqual(listing.json(), [])

        # 重复删除 → 404；非法标识 → 400（直接打函数层，避免路由层把 .. 规范化）。
        again = self.client.delete(f"/projects/{self.project_id}/exports/{export_id}", headers=self.headers)
        self.assertEqual(again.status_code, 404, again.text)
        from literature.exporters import delete_export as _delete

        from literature.storage import StorageError

        with self.assertRaises(StorageError):
            _delete(self.store, self.project_id, "../translations")
        # 源数据不受影响。
        self.assertEqual(self.store.load_document(self.project_id).blocks[2].source_text, "Hello world.")

    def test_rename_updates_manifest_and_renames_single_file(self) -> None:
        created = self._export(formats=["epub"], mode="bilingual", filename="custom name")
        self.assertEqual(created.status_code, 200, created.text)
        export_id = created.json()["export_id"]
        old_file = self.project_dir / "exports" / export_id / "custom name.epub"
        self.assertTrue(old_file.is_file())

        response = self.client.patch(
            f"/projects/{self.project_id}/exports/{export_id}", headers=self.headers,
            json={"display_name": "我的书：2026 版"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        value = response.json()
        self.assertEqual(value["display_name"], "我的书：2026 版")
        self.assertEqual([item["rel_path"] for item in value["files"]], ["我的书：2026 版.epub"])
        self.assertFalse(old_file.exists())
        new_file = self.project_dir / "exports" / export_id / "我的书：2026 版.epub"
        self.assertTrue(new_file.is_file())

        manifest = json.loads((self.project_dir / "exports" / export_id / "export.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["display_name"], "我的书：2026 版")
        listing = self.client.get(f"/projects/{self.project_id}/exports", headers=self.headers)
        self.assertEqual(listing.json()[0]["display_name"], "我的书：2026 版")

        # 重命名后的产物仍能按新路径下载。
        download = self.client.get(
            f"/projects/{self.project_id}/exports/{export_id}/file", headers=self.headers,
            params={"name": "我的书：2026 版.epub"},
        )
        self.assertEqual(download.status_code, 200, download.text)

    def test_rename_renames_all_single_files_with_collision_suffix(self) -> None:
        # markdown 导出含 assets/ 下两张同名 png：重命名后在同目录碰撞，自动加 " (2)" 后缀。
        created = self._export(formats=["markdown"], mode="source")
        self.assertEqual(created.status_code, 200, created.text)
        export_id = created.json()["export_id"]
        response = self.client.patch(
            f"/projects/{self.project_id}/exports/{export_id}", headers=self.headers,
            json={"display_name": "同名"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        rels = sorted(item["rel_path"] for item in response.json()["files"])
        self.assertEqual(rels, ["assets/同名 (2).png", "assets/同名.png", "同名.md"])
        for rel in rels:
            self.assertTrue((self.project_dir / "exports" / export_id / rel).is_file(), rel)

    def test_rename_empty_name_rejected(self) -> None:
        created = self._export(formats=["markdown"], mode="source")
        export_id = created.json()["export_id"]
        response = self.client.patch(
            f"/projects/{self.project_id}/exports/{export_id}", headers=self.headers,
            json={"display_name": "   "},
        )
        self.assertEqual(response.status_code, 400, response.text)


class FootnoteExportTests(unittest.TestCase):
    """页下注 → 真脚注：EPUB noteref/aside 与 DOCX word/footnotes.xml。

    匹配规则：正文 ``^{N}`` 上标标记 ↔ 注文 ``^{N} …`` / ``N. …``；
    同号重复引用同脚注；无编号或无对应标记的注文保持原样渲染。
    """

    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        source = self.root / "source.html"
        source.write_text("<h1>Book</h1>", encoding="utf-8")
        self.store = ProjectStore(self.root / "library")
        created = self.store.create_project(source, title="Book")
        self.project_id = created["project_id"]
        self.project_dir = Path(created["path"])
        blocks = [
            Block("ft1", "title", 0, 0, source_text="Chapter One", level=1),
            Block("fp1", "paragraph", 1, 0, source_text="Body one^{1} continues^{2}."),
            Block("fp2", "paragraph", 2, 0, source_text="Body two^{1} again."),
            Block("fn1", "footnote", 3, 0, source_text="1. First note."),
            Block("fn2", "footnote", 4, 0, source_text="^{2} Second note."),
            Block("fn3", "footnote", 5, 0, source_text="Running head without number."),
            Block("ft2", "title", 6, 0, source_text="Chapter Two", level=1),
            Block("fp3", "paragraph", 7, 0, source_text="Second chapter^{9} end."),
            Block("fp4", "paragraph", 8, 0, source_text="Stray^{7} marker without note."),
            Block("fn9", "footnote", 9, 0, source_text="^{9} Ninth note."),
            Block("fn8", "footnote", 10, 0, source_text="8. Never referenced."),
        ]
        self.document = Document(
            project_id=self.project_id,
            title="Book",
            source_language="en",
            target_language="zh-CN",
            blocks=blocks,
            toc=[TocNode("n1", "Chapter One", 1, "ft1", 0, []), TocNode("n2", "Chapter Two", 1, "ft2", 0, [])],
        )
        self.store.save_document(self.document)
        self.store.save_translation(TranslationLayer(
            project_id=self.project_id,
            locale="zh-CN",
            entries={
                "fp1": TranslationEntry("fp1", text="正文一^{1}继续^{2}。"),
                "fp2": TranslationEntry("fp2", text="正文二^{1}。"),
                "fp3": TranslationEntry("fp3", text="第二章^{9}结束。"),
                "fn1": TranslationEntry("fn1", text="第一条注。"),
                "fn2": TranslationEntry("fn2", text="^{2} 第二条注。"),
                "fn9": TranslationEntry("fn9", text="第九条注。"),
            },
        ))

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _export(self, **payload):
        return create_export(self.store, self.project_id, **payload)

    def test_plan_binds_markers_to_notes_and_repeats_share_one_note(self) -> None:
        plan = build_plan(self.document.blocks)
        self.assertEqual(set(plan.links), {"fn1", "fn2", "fn9", "fn8"})
        self.assertEqual(plan.links["fn1"].number, "1")
        self.assertEqual([ref.host_block_id for ref in plan.links["fn1"].refs], ["fp1", "fp2"])
        self.assertEqual(plan.links["fn2"].number, "2")
        self.assertEqual(plan.links["fn2"].text, "Second note.")
        self.assertEqual(plan.links["fn9"].number, "9")
        # fn8 有编号但正文无对应标记（MinerU 丢上标的场景）：回退锚定到最近正文块末尾。
        fallback = plan.links["fn8"].refs[0]
        self.assertTrue(fallback.is_fallback)
        self.assertEqual(fallback.host_block_id, "fp4")
        self.assertEqual(plan.links["fn8"].body, "Never referenced.")
        # fn3 无编号且上一条注文以句号收尾：不算续行，保持页末注原样。
        self.assertNotIn("fn3", plan.links)
        # ^{7} 没有对应注文，不生成任何引用。
        self.assertEqual([ref.number for ref in plan.refs_by_host["fp4"]], ["8"])

    def test_epub_converts_linked_notes_to_asides(self) -> None:
        value = self._export(formats=["epub"], mode="source")
        path = Path(value["path"]) / "Book.epub"
        validate_epub(path)
        with zipfile.ZipFile(path) as archive:
            bodies = {name: archive.read(name).decode("utf-8")
                      for name in archive.namelist() if name.startswith("EPUB/text/")}
        text = "".join(bodies.values())
        self.assertIn('<a id="ref-fn-fn1-1" class="noteref" epub:type="noteref" href="#fn-fn1"><sup>1</sup></a>', text)
        self.assertIn('<a id="ref-fn-fn2-1" class="noteref" epub:type="noteref" href="#fn-fn2"><sup>2</sup></a>', text)
        self.assertIn('<aside class="footnotes" epub:type="footnote" id="fn-fn1">', text)
        self.assertIn("First note.", text)
        self.assertIn("Second note.", text)
        # 注文文本不再含引导标记；全文只出现一次（aside 内），正文流不重复。
        self.assertNotIn("^{2} Second note.", text)
        self.assertEqual(text.count("1. First note."), 1)
        # 回退锚定：fp4 末尾补回编号 8 的上标引用（印刷书里有、MinerU 丢了）。
        self.assertIn("Stray<sup>7</sup> marker without note."
                      '<a id="ref-fn-fn8-1" class="noteref" epub:type="noteref" href="#fn-fn8">'
                      "<sup>8</sup></a>", text)
        # 注文统一在章节末尾：fn9/fn8 的 aside 排在第二章正文之后，fn8 不再是页末注段落。
        self.assertLess(text.find("Stray<sup>7</sup>"), text.find('id="fn-fn9"'))
        self.assertLess(text.find('id="fn-fn9"'), text.find('id="fn-fn8"'))
        # fn8 的注文只出现一次，且位于 fn8 aside 内（章节末尾），紧跟 ↩ 回链。
        self.assertEqual(text.count('<p class="footnote source">8. Never referenced.</p>'), 1)
        self.assertIn('<p class="footnote source">8. Never referenced.</p>'
                      '</section><p class="noteref-back">'
                      '<a epub:type="footnote-return" href="#ref-fn-fn8-1">\u21a9</a></p>', text)
        # 无编号注（上一条以句号收尾）保持页末注段落。
        self.assertIn('<p class="footnote source">Running head without number.</p>', text)
        self.assertNotIn("^{", text)
        # 绑定注文的正文标记不再是字面量。
        self.assertNotIn("^{9}", text)

    def test_epub_bilingual_aside_has_source_and_translation(self) -> None:
        value = self._export(formats=["epub"], mode="bilingual")
        path = Path(value["path"]) / "Book.epub"
        validate_epub(path)
        with zipfile.ZipFile(path) as archive:
            text = "".join(archive.read(name).decode("utf-8")
                           for name in archive.namelist() if name.startswith("EPUB/text/"))
        self.assertIn('<aside class="footnotes" epub:type="footnote" id="fn-fn1">', text)
        self.assertIn('<p class="footnote source" xml:lang="en" lang="en">1. First note.</p>', text)
        self.assertIn('<p class="footnote translation">第一条注。</p>', text)
        # 注文自身的引导标记在译文里同样不渲染（``^{2} 第二条注。`` → 正文）。
        self.assertIn('<p class="footnote translation">第二条注。</p>', text)
        # 原文 + 译文角色里的同号标记都转成引用（fp1 2+2、fp2 1+1、fp3 1+1），
        # 另加 fn8 的回退锚点（fp4 原文段末尾）。
        self.assertEqual(text.count('epub:type="noteref"'), 9)

    def test_epub_renders_plain_superscripts(self) -> None:
        """未绑到脚注的 ``^{…}``（指数/分数/序数）按上标渲染，不当 LaTeX 留在成品里。"""
        self.document.blocks.append(Block("fp5", "paragraph", 11, 0, source_text="Area 3^{2} m and 1/2^{1/2} and the 4^{th} line."))
        self.store.save_document(self.document)
        value = self._export(formats=["epub"], mode="source")
        path = Path(value["path"]) / "Book.epub"
        with zipfile.ZipFile(path) as archive:
            text = "".join(archive.read(name).decode("utf-8")
                           for name in archive.namelist() if name.startswith("EPUB/text/"))
        self.assertIn("Area 3<sup>2</sup> m and 1/2<sup>1/2</sup> and the 4<sup>th</sup> line.", text)
        self.assertNotIn("^{", text)

    def test_docx_writes_real_footnotes(self) -> None:
        value = self._export(formats=["docx"], mode="bilingual")
        path = Path(value["path"]) / "Book.docx"
        self.assertTrue(zipfile.is_zipfile(path))
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            self.assertIn("word/footnotes.xml", names)
            footnotes = archive.read("word/footnotes.xml").decode("utf-8")
            document_xml = archive.read("word/document.xml").decode("utf-8")
            content_types = archive.read("[Content_Types].xml").decode("utf-8")
        self.assertIn("footnotes+xml", content_types)
        self.assertIn("w:footnoteReference", document_xml)
        self.assertIn("First note.", footnotes)
        self.assertIn("Second note.", footnotes)
        self.assertIn("第一条注。", footnotes)
        self.assertNotIn("^{2} Second", footnotes)
        # 未编号注（上一条以句号收尾）仍是普通段落；已绑定注文不再出现在正文流。
        self.assertIn("Running head without number.", document_xml)
        # fn8 经页位回退锚定进入真脚注，正文流不再保留。
        self.assertIn("Never referenced.", footnotes)
        self.assertNotIn("Never referenced.", document_xml)
        self.assertNotIn("First note.", document_xml)
        # 引用数：fn1 首锚、fn2、fn9、fn8 回退锦各一次（重复引用降级为上标文字）。
        self.assertEqual(document_xml.count("<w:footnoteReference"), 4)
        # 未绑定标记同样降级为上标 run，文档里不留 ``^{…}``。
        self.assertNotIn("^{", document_xml)
        self.assertNotIn("^{", footnotes)
        self.assertIn('w:vertAlign w:val="superscript"', document_xml)
        from docx import Document as WordDocument

        word = WordDocument(str(path))  # 能正常打开
        self.assertTrue(word.paragraphs)


    def test_plan_merges_continuations_and_parses_more_leading_numbers(self) -> None:
        """无编号续行并入上一条；Unicode 上标、``N␣`` 前缀也能识别。"""
        blocks = [
            Block("t0", "title", 0, 0, source_text="C", level=1),
            Block("p0", "paragraph", 1, 0, source_text="Text.^{1}"),
            Block("fa", "footnote", 2, 0, source_text="1. See the chart below"),
            Block("fb", "footnote", 3, 0, source_text="which was redrawn for this edition."),
            Block("fc", "footnote", 4, 0, source_text="\u00b9 Unicode leading."),
            Block("fd", "footnote", 5, 0, source_text="15 Bare leading."),
        ]
        plan = build_plan(blocks)
        self.assertEqual(set(plan.absorbed), {"fb"})
        self.assertEqual(plan.links["fa"].text, "1. See the chart below which was redrawn for this edition.")
        self.assertEqual(plan.links["fa"].number, "1")
        self.assertEqual(plan.links["fc"].number, "1")
        self.assertEqual(plan.links["fd"].number, "15")
        # 三条注都经页位回退（或精确）锦到正文块末尾。
        self.assertEqual(plan.links["fa"].refs[0].host_block_id, "p0")
        self.assertEqual(plan.links["fc"].refs[0].host_block_id, "p0")
        self.assertEqual(plan.links["fd"].refs[0].host_block_id, "p0")
        self.assertEqual(plan.links["fd"].body, "Bare leading.")

    def test_epub_asides_live_at_section_end(self) -> None:
        """aside 顺序 = 锚点顺序，全部在章节正文之后、nav 之外。"""
        value = self._export(formats=["epub"], mode="source")
        path = Path(value["path"]) / "Book.epub"
        with zipfile.ZipFile(path) as archive:
            chapter_two = archive.read("EPUB/text/section-0002.xhtml").decode("utf-8")
        body_end = chapter_two.rindex("</body>")
        for note in ("fn-fn9", "fn-fn8"):
            anchor = chapter_two.index(f'id="{note}"')
            self.assertLess(anchor, body_end)
        self.assertGreater(chapter_two.index('id="fn-fn8"'), chapter_two.index("Stray<sup>7</sup>"))

    def test_markdown_keeps_page_note_rendering(self) -> None:
        value = self._export(formats=["markdown"], mode="source")
        markdown = (Path(value["path"]) / "Book.md").read_text(encoding="utf-8")
        self.assertIn("> [页下注] 1. First note.", markdown)
        self.assertIn("^{1}", markdown)


if __name__ == "__main__":
    unittest.main()
